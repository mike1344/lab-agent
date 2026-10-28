"""核心业务接口回归测试，使用独立 SQLite 数据库，不会改动开发 MySQL 数据。"""

import base64
import gc
import json
import os
import tempfile
import unittest
from datetime import date, datetime, time, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

# 测试库只在系统临时目录中创建，不删除或覆盖项目内的同名文件。
TEST_DIRECTORY = tempfile.TemporaryDirectory(prefix="lab_agent_api_test_", ignore_cleanup_errors=True)
TEST_DB = Path(TEST_DIRECTORY.name) / "test_lab_agent.db"

# 必须在导入 app 之前覆盖数据库配置，因为 engine 会在导入阶段创建。
os.environ["DATABASE_URL"] = f"sqlite+pysqlite:///{TEST_DB.as_posix()}"
os.environ["OPENAI_API_KEY"] = ""
# 单元测试里直接调用定时任务的纯函数，不让应用后台线程并发跑调度任务干扰断言。
os.environ["SCHEDULER_ENABLED"] = "false"
# 登录安全基线：常规用例关闭图形验证码；验证码与锁定行为由专门用例覆盖。
os.environ["LOGIN_CAPTCHA_ENABLED"] = "false"
# 邮件通道基线：测试环境不配置 SMTP，发送逻辑应静默跳过（不连网）。
os.environ["SMTP_HOST"] = ""
os.environ["SMTP_FROM"] = ""
# 限流基线：常规用例关闭；限流行为由专门用例临时打开 settings.RATE_LIMIT_ENABLED 覆盖。
os.environ["RATE_LIMIT_ENABLED"] = "false"
# 电子签名基线：常规用例关闭密码二次确认；专门用例临时打开 settings.SIGNATURE_REQUIRED 覆盖。
os.environ["SIGNATURE_REQUIRED"] = "false"
# 二次验证有效期基线：测试里逐次验证，避免 TTL 缓存让用例互相影响；TTL 行为由专门用例覆盖。
os.environ["SIGNATURE_TTL_MINUTES"] = "0"
# 病毒扫描基线：测试环境不配置扫描器，上传直接放行（扫描逻辑由专门用例用 mock 覆盖）。
os.environ["CLAMAV_COMMAND"] = ""

from app.ai.agent import execute_tools_node, has_explicit_reservation_intent, local_tool_node
from app.ai.knowledge import build_local_embedding
from app.database import SessionLocal, engine
from app.main import app
from app.models.reservation import Reservation
from app.models.user import User
from app.utils.password import hash_password
from chromadb.api.client import SharedSystemClient
from fastapi.testclient import TestClient


def extract_streamed_tokens(raw_text: str) -> str:
    """把 SSE 响应里的 token 事件拼接成完整回答，避免断言依赖分段边界。"""
    content = []
    for line in raw_text.splitlines():
        if line.startswith("data: "):
            payload = json.loads(line[len("data: ") :])
            if "content" in payload:
                content.append(payload["content"])
    return "".join(content)


class ApiFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        # AI SSE 会调用真实 RAG；索引也放在临时目录，避免测试改动开发知识库。
        cls.chroma_directory = tempfile.TemporaryDirectory(prefix="lab_agent_chroma_test_", ignore_cleanup_errors=True)
        cls.chroma_patch = patch("app.ai.knowledge.CHROMA_DIR", Path(cls.chroma_directory.name))
        cls.chroma_patch.start()
        with SessionLocal() as db:
            db.add(
                User(
                    username="admin",
                    password=hash_password("Admin123456"),
                    email="admin@example.com",
                    full_name="系统管理员",
                    is_active=True,
                    is_superuser=True,
                )
            )
            db.commit()

    @classmethod
    def tearDownClass(cls):
        engine.dispose()
        cls.chroma_patch.stop()
        SharedSystemClient.clear_system_cache()
        gc.collect()
        cls.chroma_directory.cleanup()
        TEST_DIRECTORY.cleanup()

    def login(self, username: str, password: str) -> dict[str, str]:
        response = self.client.post(
            "/api/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.json()["code"], 200)
        return {"Authorization": f"Bearer {response.json()['data']['token']}"}

    def test_health_endpoint(self):
        """部署探活：/health 必须返回 200 和 status=ok，供监控系统探测。"""
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok"})

    def test_local_embedding_is_stable(self):
        """本地向量必须固定且归一化，否则同一问题每次检索结果会漂移。"""
        first_vector = build_local_embedding("实验室设备故障")
        second_vector = build_local_embedding("实验室设备故障")

        self.assertEqual(first_vector, second_vector)
        self.assertEqual(len(first_vector), 128)
        vector_length = sum(value * value for value in first_vector) ** 0.5
        self.assertAlmostEqual(vector_length, 1.0)

    def test_operation_log_records_writes_and_guards_access(self):
        """操作日志：写操作落库（成功/失败都记），读操作不记录，仅管理员可查询。"""
        from app.models.operation_log import OperationLog

        admin_headers = self.login("admin", "Admin123456")
        student_headers = self.login("student01", "Student123")

        with SessionLocal() as db:
            before = db.query(OperationLog).count()

        # 1) 读操作不产生审计噪音
        self.assertEqual(self.client.get("/api/labs", headers=admin_headers).status_code, 200)
        with SessionLocal() as db:
            self.assertEqual(db.query(OperationLog).count(), before)

        # 2) 成功的写操作落库（管理员新增用户）
        created = self.client.post(
            "/api/user",
            headers=admin_headers,
            json={"username": "audit_user", "password": "Audit123456", "email": "audit_user@example.com"},
        ).json()
        self.assertEqual(created["code"], 200)

        # 3) 失败的写操作也落库（用户名重复 → 400）
        failed = self.client.post(
            "/api/user",
            headers=admin_headers,
            json={"username": "audit_user", "password": "Audit123456", "email": "audit_other@example.com"},
        )
        self.assertEqual(failed.status_code, 400)

        with SessionLocal() as db:
            logs = db.query(OperationLog).order_by(OperationLog.id.desc()).limit(10).all()

        success_log = next(log for log in logs if log.path == "/api/user" and log.success)
        self.assertEqual(success_log.method, "POST")
        self.assertEqual(success_log.username, "admin")
        self.assertTrue(success_log.is_superuser)
        self.assertEqual(success_log.http_status, 200)
        # 操作说明取自路由 docstring 首行（中文字段，便于审计阅读）
        self.assertIn("管理员", success_log.summary)
        self.assertGreaterEqual(success_log.duration_ms, 0)
        self.assertTrue(success_log.ip)

        failed_log = next(log for log in logs if log.path == "/api/user" and not log.success)
        self.assertEqual(failed_log.http_status, 400)
        self.assertIn("已存在", failed_log.message)

        # 4) 查询接口本身不写日志，且仅管理员可用
        self.assertEqual(self.client.get("/api/operation-logs", headers=student_headers).status_code, 403)
        self.assertEqual(
            self.client.get("/api/operation-logs/tool-calls", headers=student_headers).status_code,
            403,
        )
        tool_calls = self.client.get("/api/operation-logs/tool-calls", headers=admin_headers).json()
        self.assertEqual(tool_calls["code"], 200)
        res = self.client.get(
            "/api/operation-logs",
            headers=admin_headers,
            params={"username": "admin", "method": "post", "success": "true"},
        ).json()
        self.assertEqual(res["code"], 200)
        payload = res["data"]
        self.assertGreaterEqual(payload["total"], 1)
        self.assertTrue(all(item["username"] == "admin" for item in payload["list"]))
        self.assertTrue(all(item["method"] == "POST" for item in payload["list"]))
        self.assertTrue(all(item["success"] for item in payload["list"]))

    def test_announcement_comments_and_read_receipts(self):
        """公告评论/回复与已读回执：幂等已读、二级回复、删除权限、管理员回执统计。"""
        admin_headers = self.login("admin", "Admin123456")
        # 本用例按字母序会先于 test_full_business_flow 执行，因此自带一个普通用户
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "ann_commenter",
                    "password": "AnnCom123456",
                    "email": "ann_commenter@example.com",
                    "full_name": "公告评论员",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("ann_commenter", "AnnCom123456")

        created = self.client.post(
            "/api/announcements",
            headers=admin_headers,
            json={"title": "实验楼停电通知", "content": "本周六 8:00-12:00 实验楼停电检修。", "is_pinned": True},
        ).json()
        self.assertEqual(created["code"], 200)
        ann_id = created["data"]["id"]

        # 1) 初始未读；标记已读后列表同步，重复标记幂等（不会产生第二条回执）
        listed = self.client.get("/api/announcements", headers=student_headers).json()["data"]
        row = next(item for item in listed if item["id"] == ann_id)
        self.assertFalse(row["is_read"])
        self.assertEqual(row["read_count"], 0)

        self.assertEqual(
            self.client.post(f"/api/announcements/{ann_id}/read", headers=student_headers).json()["data"]["read_count"],
            1,
        )
        self.assertEqual(
            self.client.post(f"/api/announcements/{ann_id}/read", headers=student_headers).json()["data"]["read_count"],
            1,
        )
        listed = self.client.get("/api/announcements", headers=student_headers).json()["data"]
        row = next(item for item in listed if item["id"] == ann_id)
        self.assertTrue(row["is_read"])
        self.assertEqual(row["read_count"], 1)

        # 2) 学生发一级评论，管理员回复；列表按时间正序返回并带用户名
        top = self.client.post(
            f"/api/announcements/{ann_id}/comments",
            headers=student_headers,
            json={"content": "  停电期间设备需要断电吗？  "},
        ).json()["data"]
        self.assertEqual(top["content"], "停电期间设备需要断电吗？")
        self.assertIsNone(top["parent_id"])
        self.assertTrue(top["username"])

        reply = self.client.post(
            f"/api/announcements/{ann_id}/comments",
            headers=admin_headers,
            json={"content": "需要，请提前关闭并拔掉电源。", "parent_id": top["id"]},
        ).json()["data"]
        self.assertEqual(reply["parent_id"], top["id"])

        comments = self.client.get(f"/api/announcements/{ann_id}/comments", headers=student_headers).json()["data"]
        self.assertEqual([item["id"] for item in comments], [top["id"], reply["id"]])

        # 3) 不支持回复"回复"，只能一层
        nested = self.client.post(
            f"/api/announcements/{ann_id}/comments",
            headers=student_headers,
            json={"content": "再回复一层", "parent_id": reply["id"]},
        )
        self.assertEqual(nested.json()["code"], 400)

        # 4) 只有作者或管理员能删评论
        self.assertEqual(
            self.client.delete(f"/api/announcements/comments/{reply['id']}", headers=student_headers).status_code,
            403,
        )
        # 管理员删除一级评论会连带删除其回复
        self.assertEqual(
            self.client.delete(f"/api/announcements/comments/{top['id']}", headers=admin_headers).status_code,
            200,
        )
        self.assertEqual(
            self.client.get(f"/api/announcements/{ann_id}/comments", headers=student_headers).json()["data"],
            [],
        )

        # 5) 已读回执：管理员可查已读人数/在册人数/名单；学生无权查看
        self.assertEqual(
            self.client.get(f"/api/announcements/{ann_id}/read-receipts", headers=student_headers).status_code,
            403,
        )
        receipts = self.client.get(f"/api/announcements/{ann_id}/read-receipts", headers=admin_headers).json()["data"]
        self.assertEqual(receipts["read_count"], 1)
        self.assertGreaterEqual(receipts["total_users"], 2)
        self.assertEqual(len(receipts["readers"]), 1)
        self.assertTrue(receipts["readers"][0]["username"])

        # 6) 删除公告时评论与回执一并清理（应用层级联）
        self.assertEqual(self.client.delete(f"/api/announcements/{ann_id}", headers=admin_headers).status_code, 200)
        self.assertEqual(self.client.get(f"/api/announcements/{ann_id}", headers=admin_headers).status_code, 404)

    def test_favorites_and_favorite_only_filters(self):
        """收藏：幂等新增、类型过滤、列表 favorite_only 过滤、取消收藏与越权保护。"""
        admin_headers = self.login("admin", "Admin123456")
        # 本用例按字母序较早执行，用自建的实验室/设备验证，不依赖其他用例数据
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "收藏测试实验室", "location": "实训楼 601", "capacity": 8},
        ).json()
        self.assertEqual(lab["code"], 200)
        lab_id = lab["data"]["id"]
        equipment = self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={"lab_id": lab_id, "name": "收藏测试设备", "code": "FAV-001", "quantity": 1},
        ).json()
        self.assertEqual(equipment["code"], 200)
        equipment_id = equipment["data"]["id"]

        # 普通用户自建（不依赖其他用例注册的账号）
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "fav_user",
                    "password": "FavUser123456",
                    "email": "fav_user@example.com",
                },
            ).json()["code"],
            200,
        )
        fav_headers = self.login("fav_user", "FavUser123456")

        # 1) 收藏实验室：重复收藏幂等（返回同一条记录）
        first = self.client.post(
            "/api/favorites", headers=fav_headers, json={"target_type": "lab", "target_id": lab_id}
        ).json()
        self.assertEqual(first["code"], 200)
        self.assertEqual(first["data"]["name"], "收藏测试实验室")
        favorite_id = first["data"]["id"]
        again = self.client.post(
            "/api/favorites", headers=fav_headers, json={"target_type": "lab", "target_id": lab_id}
        ).json()
        self.assertEqual(again["data"]["id"], favorite_id)

        # 2) 收藏设备；不存在的目标被拒绝
        self.assertEqual(
            self.client.post(
                "/api/favorites",
                headers=fav_headers,
                json={"target_type": "equipment", "target_id": equipment_id},
            ).json()["code"],
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/favorites", headers=fav_headers, json={"target_type": "lab", "target_id": 999999}
            ).status_code,
            400,
        )

        # 3) 收藏列表：按类型过滤 + 目标名称回填
        labs = self.client.get("/api/favorites", headers=fav_headers, params={"target_type": "lab"}).json()["data"]
        self.assertEqual(len(labs), 1)
        self.assertEqual(labs[0]["target_id"], lab_id)
        equipments = self.client.get("/api/favorites", headers=fav_headers, params={"target_type": "equipment"}).json()[
            "data"
        ]
        self.assertEqual(equipments[0]["name"], "收藏测试设备")

        # 4) 列表接口的 favorite_only 过滤走数据库子查询
        filtered = self.client.get("/api/labs", headers=fav_headers, params={"favorite_only": True}).json()["data"]
        self.assertEqual([item["id"] for item in filtered], [lab_id])
        filtered_eq = self.client.get("/api/equipment", headers=fav_headers, params={"favorite_only": True}).json()[
            "data"
        ]
        self.assertEqual([item["id"] for item in filtered_eq], [equipment_id])

        # 5) 别人的收藏不能被删除（按不存在处理）
        self.assertEqual(self.client.delete(f"/api/favorites/{favorite_id}", headers=admin_headers).status_code, 404)

        # 6) 取消收藏后过滤结果为空
        self.assertEqual(self.client.delete(f"/api/favorites/{favorite_id}", headers=fav_headers).status_code, 200)
        filtered = self.client.get("/api/labs", headers=fav_headers, params={"favorite_only": True}).json()["data"]
        self.assertEqual(filtered, [])

    def test_procurement_flow_updates_stock_once(self):
        """采购记录：状态机 pending→arrived→accepted，验收只入库一次，已验收不可删除。"""
        admin_headers = self.login("admin", "Admin123456")

        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "采购测试实验室", "location": "实训楼 602", "capacity": 6},
        ).json()["data"]
        equipment = self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={"lab_id": lab["id"], "name": "采购测试设备", "code": "PROC-001", "quantity": 1},
        ).json()["data"]
        equipment_id = equipment["id"]

        created = self.client.post(
            "/api/procurements",
            headers=admin_headers,
            json={"name": "采购测试设备", "quantity": 3, "equipment_id": equipment_id, "remark": "教学用"},
        ).json()
        self.assertEqual(created["code"], 200)
        item = created["data"]
        self.assertEqual(item["status"], "pending")
        self.assertEqual(item["quantity"], 3)
        procurement_id = item["id"]

        # 1) 不能跳级：pending 直接验收被拒绝
        self.assertEqual(
            self.client.put(
                f"/api/procurements/{procurement_id}", headers=admin_headers, json={"status": "accepted"}
            ).status_code,
            400,
        )

        # 2) 到货
        arrived = self.client.put(
            f"/api/procurements/{procurement_id}", headers=admin_headers, json={"status": "arrived"}
        ).json()
        self.assertEqual(arrived["data"]["status"], "arrived")

        # 3) 验收：数量计入设备库存
        stock_before = next(
            e
            for e in self.client.get("/api/equipment", headers=admin_headers, params={"lab_id": lab["id"]}).json()[
                "data"
            ]
            if e["id"] == equipment_id
        )["quantity"]
        accepted = self.client.put(
            f"/api/procurements/{procurement_id}",
            headers=admin_headers,
            json={"status": "accepted", "remark": "验收合格"},
        ).json()
        self.assertEqual(accepted["data"]["status"], "accepted")
        stock_after = next(
            e
            for e in self.client.get("/api/equipment", headers=admin_headers, params={"lab_id": lab["id"]}).json()[
                "data"
            ]
            if e["id"] == equipment_id
        )["quantity"]
        self.assertEqual(stock_after, stock_before + 3)

        # 4) 重复验收被拒绝：不会二次入库
        self.assertEqual(
            self.client.put(
                f"/api/procurements/{procurement_id}", headers=admin_headers, json={"status": "accepted"}
            ).status_code,
            400,
        )

        # 5) 已验收记录不可删除；列表回填设备名与申请人
        self.assertEqual(
            self.client.delete(f"/api/procurements/{procurement_id}", headers=admin_headers).status_code, 400
        )
        listed = self.client.get(
            "/api/procurements",
            headers=admin_headers,
            params={"status": "accepted", "keyword": "采购测试"},
        ).json()["data"]
        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0]["equipment_name"], "采购测试设备")
        self.assertEqual(listed[0]["applicant_name"], "系统管理员")

        # 6) 普通用户无权访问采购接口
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "proc_student",
                    "password": "ProcUser123456",
                    "email": "proc_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("proc_student", "ProcUser123456")
        self.assertEqual(self.client.get("/api/procurements", headers=student_headers).status_code, 403)

        # 7) pending 记录允许删除
        pending = self.client.post(
            "/api/procurements", headers=admin_headers, json={"name": "临时采购", "quantity": 1}
        ).json()["data"]
        self.assertEqual(
            self.client.delete(f"/api/procurements/{pending['id']}", headers=admin_headers).status_code, 200
        )

    def test_signature_ttl_allows_repeat_within_window(self):
        """二次验证有效期：首次必须输入密码；TTL 窗口内重复危险操作免输；过期后重新要求。"""
        from app.config import settings
        from app.dependencies.signature import reset_signature_cache

        admin_headers = self.login("admin", "Admin123456")
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "TTL 验证实验室", "location": "实训楼 904", "capacity": 5},
        ).json()["data"]

        original_required = settings.SIGNATURE_REQUIRED
        original_ttl = settings.SIGNATURE_TTL_MINUTES
        settings.SIGNATURE_REQUIRED = True
        settings.SIGNATURE_TTL_MINUTES = 5
        reset_signature_cache()
        try:
            # 1) 首次危险操作：不带密码被拒绝
            first = self.client.post(
                "/api/equipment",
                headers=admin_headers,
                json={"lab_id": lab["id"], "name": "TTL 设备一", "code": "TTL-001", "quantity": 1},
            ).json()["data"]
            denied = self.client.delete(f"/api/equipment/{first['id']}", headers=admin_headers)
            self.assertEqual(denied.status_code, 400)

            # 2) 正确密码通过，并进入 5 分钟免密窗口
            ok = self.client.delete(
                f"/api/equipment/{first['id']}",
                headers=admin_headers,
                params={"signature_password": "Admin123456"},
            )
            self.assertEqual(ok.status_code, 200)

            # 3) TTL 内第二次危险操作无需密码
            second = self.client.post(
                "/api/equipment",
                headers=admin_headers,
                json={"lab_id": lab["id"], "name": "TTL 设备二", "code": "TTL-002", "quantity": 1},
            ).json()["data"]
            allowed = self.client.delete(f"/api/equipment/{second['id']}", headers=admin_headers)
            self.assertEqual(allowed.status_code, 200)

            # 4) 缓存清空（等价于过期）后重新要求密码
            reset_signature_cache()
            third = self.client.post(
                "/api/equipment",
                headers=admin_headers,
                json={"lab_id": lab["id"], "name": "TTL 设备三", "code": "TTL-003", "quantity": 1},
            ).json()["data"]
            expired = self.client.delete(f"/api/equipment/{third['id']}", headers=admin_headers)
            self.assertEqual(expired.status_code, 400)
        finally:
            settings.SIGNATURE_REQUIRED = original_required
            settings.SIGNATURE_TTL_MINUTES = original_ttl
            reset_signature_cache()

    def test_usage_report_and_export(self):
        """报表中心：机时统计按实验室/用户汇总，费用按单价预估，导出双工作表。"""
        from datetime import date, timedelta
        from io import BytesIO

        from openpyxl import load_workbook

        admin_headers = self.login("admin", "Admin123456")
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "报表测试实验室", "location": "实训楼 903", "capacity": 20},
        ).json()["data"]
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "report_student",
                    "password": "ReportStu123456",
                    "email": "report_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("report_student", "ReportStu123456")

        target = date.today() + timedelta(days=1)
        reservation_ids = []
        for start, end in (("08:00:00", "09:00:00"), ("10:00:00", "11:00:00")):
            created = self.client.post(
                "/api/reservations",
                headers=student_headers,
                json={
                    "lab_id": lab["id"],
                    "reservation_date": target.isoformat(),
                    "start_time": start,
                    "end_time": end,
                    "purpose": "报表统计",
                },
            ).json()
            self.assertEqual(created["code"], 200)
            reservation_ids.append(created["data"][0]["id"])
            self.assertEqual(
                self.client.put(
                    f"/api/reservations/{reservation_ids[-1]}/review",
                    headers=admin_headers,
                    json={"status": "approved", "review_comment": "同意"},
                ).json()["code"],
                200,
            )

        report = self.client.get(
            "/api/reports/usage",
            headers=admin_headers,
            params={"start_date": target.isoformat(), "end_date": target.isoformat(), "lab_id": lab["id"]},
        ).json()["data"]
        self.assertEqual(report["totals"]["approved"], 2)
        self.assertAlmostEqual(report["totals"]["hours"], 2.0)
        self.assertEqual(report["totals"]["users"], 1)
        self.assertAlmostEqual(report["totals"]["estimated_cost"], 2.0 * report["hour_rate"])
        self.assertEqual(report["by_lab"][0]["lab_name"], "报表测试实验室")
        self.assertEqual(report["by_lab"][0]["users"], 1)
        self.assertEqual(report["by_user"][0]["username"], "report_student")

        # 导出：xlsx 内含「按实验室」「按用户」两个工作表
        exported = self.client.get(
            "/api/reports/usage/export",
            headers=admin_headers,
            params={"start_date": target.isoformat(), "end_date": target.isoformat()},
        )
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content))
        self.assertEqual(workbook.sheetnames, ["按实验室", "按用户"])

        # 普通用户不能看报表
        self.assertEqual(self.client.get("/api/reports/usage", headers=student_headers).status_code, 403)

    def test_audit_retention_signature_and_upload_scan(self):
        """合规三件套：审计留存清理、危险操作电子签名、上传病毒扫描钩子。"""
        import base64
        from datetime import date, datetime, timedelta
        from unittest.mock import patch

        from app.config import settings
        from app.models.agent_tool_log import AgentToolLog
        from app.models.login_log import LoginLog
        from app.models.operation_log import OperationLog
        from app.services.antivirus import scan_uploaded_file
        from app.services.audit import purge_expired_audit_logs

        # 1) 审计留存：超过保留期的记录被清理，保留期内的记录保留
        old_time = datetime.now() - timedelta(days=settings.AUDIT_RETENTION_DAYS + 10)
        with SessionLocal() as db:
            db.add(
                OperationLog(
                    username="audit_old",
                    is_superuser=False,
                    method="POST",
                    path="/api/labs",
                    summary="过期审计",
                    success=True,
                    http_status=200,
                    message="ok",
                    duration_ms=1,
                    create_time=old_time,
                )
            )
            db.add(LoginLog(username="audit_old", success=True, message="过期登录", create_time=old_time))
            db.add(
                OperationLog(
                    username="audit_keep",
                    is_superuser=False,
                    method="POST",
                    path="/api/labs",
                    summary="新审计",
                    success=True,
                    http_status=200,
                    message="ok",
                    duration_ms=1,
                )
            )
            db.add(LoginLog(username="audit_keep", success=True, message="新登录"))
            admin_user = db.query(User).filter(User.username == "admin").first()
            self.assertIsNotNone(admin_user)
            db.add(
                AgentToolLog(
                    user_id=admin_user.id,
                    tool_name="old_tool",
                    arguments="{}",
                    result="",
                    success=True,
                    duration_ms=1,
                    idempotency_key="audit-old-tool-key",
                    note="过期工具调用",
                    create_time=old_time,
                )
            )
            db.add(
                AgentToolLog(
                    user_id=admin_user.id,
                    tool_name="new_tool",
                    arguments="{}",
                    result="",
                    success=True,
                    duration_ms=1,
                    idempotency_key="audit-new-tool-key",
                    note="保留期内工具调用",
                )
            )
            db.commit()
            deleted = purge_expired_audit_logs(db)
            self.assertGreaterEqual(deleted, 3)
            self.assertEqual(db.query(OperationLog).filter(OperationLog.username == "audit_old").count(), 0)
            self.assertEqual(db.query(LoginLog).filter(LoginLog.username == "audit_old").count(), 0)
            self.assertEqual(db.query(OperationLog).filter(OperationLog.username == "audit_keep").count(), 1)
            self.assertEqual(db.query(AgentToolLog).filter(AgentToolLog.tool_name == "old_tool").count(), 0)
            self.assertEqual(db.query(AgentToolLog).filter(AgentToolLog.tool_name == "new_tool").count(), 1)

        # 2) 电子签名：危险操作需要当前账号密码；错误/缺失都拒绝，正确才放行
        admin_headers = self.login("admin", "Admin123456")
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "签名测试实验室", "location": "实训楼 902", "capacity": 5},
        ).json()["data"]
        equipment = self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={"lab_id": lab["id"], "name": "签名测试设备", "code": "SIG-001", "quantity": 1},
        ).json()["data"]

        original_signature = settings.SIGNATURE_REQUIRED
        settings.SIGNATURE_REQUIRED = True
        try:
            missing = self.client.delete(f"/api/equipment/{equipment['id']}", headers=admin_headers)
            self.assertEqual(missing.status_code, 400)
            self.assertIn("二次确认", missing.json()["message"])
            wrong = self.client.delete(
                f"/api/equipment/{equipment['id']}",
                headers=admin_headers,
                params={"signature_password": "WrongPassword123"},
            )
            self.assertEqual(wrong.status_code, 400)
            # 正确密码才真正删除
            ok = self.client.delete(
                f"/api/equipment/{equipment['id']}",
                headers=admin_headers,
                params={"signature_password": "Admin123456"},
            )
            self.assertEqual(ok.status_code, 200)
        finally:
            settings.SIGNATURE_REQUIRED = original_signature

        # 3) 审核驳回也要求二次确认
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "sign_student", "password": "SignStu123456", "email": "sign_student@example.com"},
            ).json()["code"],
            200,
        )
        student_headers = self.login("sign_student", "SignStu123456")
        tomorrow = date.today() + timedelta(days=1)
        reservation = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab["id"],
                "reservation_date": tomorrow.isoformat(),
                "start_time": "08:00:00",
                "end_time": "09:00:00",
                "purpose": "签名测试",
            },
        ).json()["data"][0]
        settings.SIGNATURE_REQUIRED = True
        try:
            denied = self.client.put(
                f"/api/reservations/{reservation['id']}/review",
                headers=admin_headers,
                json={"status": "rejected", "review_comment": "无签名"},
            )
            self.assertEqual(denied.status_code, 400)
            accepted = self.client.put(
                f"/api/reservations/{reservation['id']}/review",
                headers=admin_headers,
                params={"signature_password": "Admin123456"},
                json={"status": "rejected", "review_comment": "签名有效"},
            )
            self.assertEqual(accepted.status_code, 200)
        finally:
            settings.SIGNATURE_REQUIRED = original_signature

        # 4) 审计报告导出：操作日志与登录日志都能导出，普通用户 403
        for path in ("/api/operation-logs/export", "/api/auth/login-logs/export"):
            exported = self.client.get(path, headers=admin_headers)
            self.assertEqual(exported.status_code, 200)
            self.assertIn("spreadsheetml", exported.headers["content-type"])
            self.assertEqual(self.client.get(path, headers=student_headers).status_code, 403)

        # 5) 上传病毒扫描：未配置扫描器时放行并标记 skipped；扫描器判定感染时删除文件并拒绝
        self.assertEqual(scan_uploaded_file("ignored.bin"), (True, "skipped"))
        png_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        )
        clean = self.client.post(
            "/api/upload/image",
            headers=student_headers,
            files={"file": ("scan.png", png_bytes, "image/png")},
        ).json()
        self.assertEqual(clean["code"], 200)
        self.assertEqual(clean["data"]["scan"], "skipped")
        with patch("app.api.upload.scan_uploaded_file", return_value=(False, "infected")):
            infected = self.client.post(
                "/api/upload/image",
                headers=student_headers,
                files={"file": ("scan2.png", png_bytes, "image/png")},
            )
        self.assertEqual(infected.status_code, 400)
        self.assertIn("安全扫描", infected.json()["message"])

    def test_observability_and_rate_limit(self):
        """可观测性：request_id 透传、/ready 探针、/metrics 指标；限流触发 429。"""
        from app.config import settings
        from app.middleware.rate_limit import reset_rate_limits

        # 1) X-Request-ID：外部传入时原样回写，未传入时自动生成
        echoed = self.client.get("/health", headers={"X-Request-ID": "req-fixed-001"})
        self.assertEqual(echoed.headers.get("x-request-id"), "req-fixed-001")
        generated = self.client.get("/health")
        self.assertEqual(len(generated.headers.get("x-request-id", "")), 32)

        # 2) 就绪探针：测试库可达 → ready；存活探针保持简单
        ready = self.client.get("/ready")
        self.assertEqual(ready.status_code, 200)
        self.assertEqual(ready.json(), {"status": "ready", "database": "ok"})

        # 3) Prometheus 指标：默认开启，包含请求计数指标
        metrics = self.client.get("/metrics")
        self.assertEqual(metrics.status_code, 200)
        self.assertIn("http_requests_total", metrics.text)

        # 4) 限流：临时打开开关，认证接口 10 次/分钟；第 11 次返回 429
        original = settings.RATE_LIMIT_ENABLED
        settings.RATE_LIMIT_ENABLED = True
        reset_rate_limits()
        try:
            statuses = []
            for index in range(settings.RATE_LIMIT_AUTH_PER_MINUTE + 1):
                response = self.client.post(
                    "/api/auth/login",
                    json={"username": f"rl_user_{index}", "password": "WrongPassword123"},
                )
                statuses.append(response.status_code)
            self.assertEqual(statuses[-1], 429)
            self.assertTrue(all(code in (401, 429) for code in statuses))
        finally:
            settings.RATE_LIMIT_ENABLED = original
            reset_rate_limits()

    def test_refresh_token_flow_and_invalidation(self):
        """刷新令牌：登录下发、换新访问令牌、类型隔离、改密后旧刷新令牌失效。"""
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "refresh_user",
                    "password": "Refresh123456",
                    "email": "refresh_user@example.com",
                },
            ).json()["code"],
            200,
        )

        login_data = self.client.post(
            "/api/auth/login", json={"username": "refresh_user", "password": "Refresh123456"}
        ).json()
        self.assertEqual(login_data["code"], 200)
        access_token = login_data["data"]["token"]
        refresh_token = login_data["data"]["refresh_token"]
        self.assertTrue(refresh_token)

        # 1) access token 可正常访问接口
        self.assertEqual(
            self.client.get("/api/user/info", headers={"Authorization": f"Bearer {access_token}"}).status_code,
            200,
        )

        # 2) refresh 换新访问令牌（同时轮换 refresh token）
        refreshed = self.client.post("/api/auth/refresh", json={"refresh_token": refresh_token}).json()
        self.assertEqual(refreshed["code"], 200)
        new_access = refreshed["data"]["token"]
        new_refresh = refreshed["data"]["refresh_token"]
        self.assertNotEqual(new_access, access_token)
        self.assertEqual(
            self.client.get("/api/user/info", headers={"Authorization": f"Bearer {new_access}"}).status_code,
            200,
        )

        # 3) 类型隔离：access 不能当 refresh 用，refresh 不能当 access 用
        self.assertEqual(
            self.client.post("/api/auth/refresh", json={"refresh_token": new_access}).json()["code"],
            401,
        )
        self.assertEqual(
            self.client.get("/api/user/info", headers={"Authorization": f"Bearer {new_refresh}"}).status_code,
            401,
        )

        # 4) 改密后旧 refresh token 立即失效（密码版本指纹不匹配）
        changed = self.client.put(
            "/api/user/password",
            headers={"Authorization": f"Bearer {new_access}"},
            json={"old_password": "Refresh123456", "new_password": "Refresh654321"},
        )
        self.assertEqual(changed.status_code, 200)
        stale = self.client.post("/api/auth/refresh", json={"refresh_token": new_refresh}).json()
        self.assertEqual(stale["code"], 401)
        self.assertIn("密码已修改", stale["message"])

    def test_email_notifications_wire_and_skip(self):
        """邮件通知：未配置 SMTP 时静默跳过；审核/报修/开课提醒都会触发发信。"""
        from datetime import datetime, timedelta
        from unittest.mock import patch

        from app.database import SessionLocal
        from app.services.email import build_notification_email, email_configured, send_notification_email
        from app.services.schedule import remind_upcoming_reservations

        # 1) 纯函数：邮件头与正文正确；未配置 SMTP 时直接跳过（不连网）
        message = build_notification_email("student@example.com", "测试主题", "测试正文")
        self.assertEqual(message["To"], "student@example.com")
        self.assertEqual(message["Subject"], "测试主题")
        self.assertIn("测试正文", message.get_content())
        self.assertFalse(email_configured())
        self.assertFalse(send_notification_email("student@example.com", "测试主题", "测试正文"))

        # 场景准备：管理员建实验室，学生注册并预约未来 30 分钟内的时段
        admin_headers = self.login("admin", "Admin123456")
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "邮件测试实验室", "location": "实训楼 901", "capacity": 8},
        ).json()["data"]
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "email_student",
                    "password": "EmailStu123456",
                    "email": "email_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("email_student", "EmailStu123456")

        start = datetime.now() + timedelta(minutes=30)
        end = start + timedelta(hours=1)
        # 预约只支持单日时段：深夜执行用例时若跨天，把结束时间截到当天最后一秒
        if end.date() != start.date():
            end = datetime.combine(start.date(), time(23, 59, 59))
        created = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab["id"],
                "reservation_date": start.date().isoformat(),
                "start_time": start.strftime("%H:%M:%S"),
                "end_time": end.strftime("%H:%M:%S"),
                "purpose": "邮件通知测试",
            },
        ).json()
        self.assertEqual(created["code"], 200)
        reservation_id = created["data"][0]["id"]

        # 2) 审核通过 → 通过 BackgroundTasks 给预约人发信
        with patch("app.api.reservation.email_service.send_notification_email") as fake_email:
            reviewed = self.client.put(
                f"/api/reservations/{reservation_id}/review",
                headers=admin_headers,
                json={"status": "approved", "review_comment": "同意"},
            ).json()
            self.assertEqual(reviewed["code"], 200)
        fake_email.assert_called_once()
        self.assertEqual(fake_email.call_args.args[0], "email_student@example.com")
        self.assertIn("审核", fake_email.call_args.args[1])

        # 3) 定时提醒 → 对即将开始的已通过预约发信
        with patch("app.services.email.send_notification_email") as reminder_email:
            with SessionLocal() as db:
                sent = remind_upcoming_reservations(db)
        self.assertGreaterEqual(sent, 1)
        reminder_email.assert_called()
        self.assertIn("预约即将开始", reminder_email.call_args.args[1])

        # 4) 报修处理 → 给报修人发信
        repair = self.client.post(
            "/api/repairs",
            headers=student_headers,
            json={"title": "邮件测试报修", "description": "显示器黑屏"},
        ).json()
        self.assertEqual(repair["code"], 200)
        with patch("app.api.repair.email_service.send_notification_email") as repair_email:
            handled = self.client.put(
                f"/api/repairs/{repair['data']['id']}/handle",
                headers=admin_headers,
                json={"status": "resolved", "handle_comment": "已更换显示器"},
            ).json()
            self.assertEqual(handled["code"], 200)
        repair_email.assert_called_once()
        self.assertIn("报修", repair_email.call_args.args[1])

    def test_multi_entity_excel_exports(self):
        """多维导出：设备台账/报修记录/用户列表三个 Excel 接口可用且仅管理员可访问。"""
        from io import BytesIO

        from openpyxl import load_workbook

        admin_headers = self.login("admin", "Admin123456")

        # 先造一条设备与一条报修，确保导出内容可断言
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "导出测试实验室", "location": "实训楼 801", "capacity": 12},
        ).json()["data"]
        self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={"lab_id": lab["id"], "name": "导出测试设备", "code": "EXP-001", "quantity": 2},
        ).json()
        self.assertEqual(
            self.client.post(
                "/api/repairs",
                headers=admin_headers,
                json={"title": "导出测试报修", "description": "投影仪不亮", "lab_id": lab["id"]},
            ).json()["code"],
            200,
        )

        def export_rows(headers, path, params=None):
            response = self.client.get(path, headers=headers, params=params or {})
            self.assertEqual(response.status_code, 200)
            self.assertIn("spreadsheetml", response.headers["content-type"])
            return list(load_workbook(BytesIO(response.content)).active.values)

        equipment_rows = export_rows(admin_headers, "/api/equipment/export", {"keyword": "导出测试"})
        self.assertEqual(tuple(equipment_rows[0][:4]), ("设备编号", "设备名称", "所属实验室", "数量"))
        self.assertTrue(any(row[0] == "EXP-001" for row in equipment_rows[1:]))

        repair_rows = export_rows(admin_headers, "/api/repairs/export")
        self.assertIn("故障描述", repair_rows[0])
        self.assertTrue(any(row[4] == "导出测试报修" for row in repair_rows[1:]))

        user_rows = export_rows(admin_headers, "/api/user/export")
        self.assertIn("用户名", user_rows[0])
        self.assertTrue(any(row[1] == "admin" for row in user_rows[1:]))

        # 普通用户不能导出
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "export_student",
                    "password": "Export123456",
                    "email": "export_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("export_student", "Export123456")
        self.assertEqual(self.client.get("/api/user/export", headers=student_headers).status_code, 403)
        self.assertEqual(self.client.get("/api/equipment/export", headers=student_headers).status_code, 403)
        self.assertEqual(self.client.get("/api/repairs/export", headers=student_headers).status_code, 403)

    def test_course_schedule_blocks_overlapping_reservations(self):
        """实验课程：课表与自由预约双向冲突校验，选课名单可增删。"""
        from datetime import date, timedelta

        admin_headers = self.login("admin", "Admin123456")

        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "课程测试实验室", "location": "实训楼 701", "capacity": 30},
        ).json()["data"]
        lab_id = lab["id"]

        # 教师账号（管理员创建，is_teacher=True；full_name 为空 → 显示 username）
        self.assertEqual(
            self.client.post(
                "/api/user",
                headers=admin_headers,
                json={
                    "username": "course_teacher",
                    "password": "Teacher123456",
                    "email": "course_teacher@example.com",
                    "is_teacher": True,
                },
            ).json()["code"],
            200,
        )
        teacher_headers = self.login("course_teacher", "Teacher123456")

        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "course_student",
                    "password": "CourseStu123456",
                    "email": "course_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("course_student", "CourseStu123456")

        today = date.today()
        weekday = 3  # 周三；若今天就是周三则取下一周，保证是未来日期
        days_ahead = (weekday - today.isoweekday()) % 7 or 7
        target = today + timedelta(days=days_ahead)

        created = self.client.post(
            "/api/courses",
            headers=teacher_headers,
            json={
                "name": "数据结构实验",
                "lab_id": lab_id,
                "weekday": weekday,
                "start_time": "08:00:00",
                "end_time": "10:00:00",
                "start_date": today.isoformat(),
                "end_date": (today + timedelta(days=60)).isoformat(),
                "description": "每周三上午",
            },
        ).json()
        self.assertEqual(created["code"], 200)
        course_id = created["data"]["id"]
        self.assertEqual(created["data"]["teacher_name"], "course_teacher")

        # 1) 课程占用：预约 09:00-10:00 与课表重叠 → 拒绝
        blocked = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": target.isoformat(),
                "start_time": "09:00:00",
                "end_time": "10:00:00",
                "purpose": "课程冲突测试",
            },
        )
        self.assertEqual(blocked.status_code, 400)
        self.assertIn("课程", blocked.json()["message"])

        # 2) 相邻时段 10:00 开始不重叠 → 正常预约
        allowed = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": target.isoformat(),
                "start_time": "10:00:00",
                "end_time": "11:00:00",
                "purpose": "课程结束后使用",
            },
        ).json()
        self.assertEqual(allowed["code"], 200)

        # 3) 课程间冲突：同实验室同周几重叠 → 拒绝
        overlap = self.client.post(
            "/api/courses",
            headers=teacher_headers,
            json={
                "name": "编译原理实验",
                "lab_id": lab_id,
                "weekday": weekday,
                "start_time": "09:00:00",
                "end_time": "11:00:00",
                "start_date": today.isoformat(),
                "end_date": (today + timedelta(days=60)).isoformat(),
            },
        )
        self.assertEqual(overlap.status_code, 400)

        # 4) 课程与已有预约冲突：10:00-12:00 覆盖学生 10:00-11:00 的预约 → 拒绝
        reservation_clash = self.client.post(
            "/api/courses",
            headers=teacher_headers,
            json={
                "name": "编译原理实验",
                "lab_id": lab_id,
                "weekday": weekday,
                "start_time": "10:00:00",
                "end_time": "12:00:00",
                "start_date": today.isoformat(),
                "end_date": (today + timedelta(days=60)).isoformat(),
            },
        )
        self.assertEqual(reservation_clash.status_code, 400)
        self.assertIn("预约", reservation_clash.json()["message"])

        # 5) 不与任何占用重叠的课程可以创建
        adjacent = self.client.post(
            "/api/courses",
            headers=teacher_headers,
            json={
                "name": "编译原理实验",
                "lab_id": lab_id,
                "weekday": weekday,
                "start_time": "12:00:00",
                "end_time": "13:00:00",
                "start_date": today.isoformat(),
                "end_date": (today + timedelta(days=60)).isoformat(),
            },
        ).json()
        self.assertEqual(adjacent["code"], 200)

        # 6) 选课名单：教师添加学生（幂等）、查看名单、移出
        students = self.client.get("/api/courses/available-students", headers=teacher_headers).json()["data"]
        student_id = next(item["id"] for item in students if item["username"] == "course_student")

        added = self.client.post(
            f"/api/courses/{course_id}/students",
            headers=teacher_headers,
            json={"student_ids": [student_id]},
        ).json()
        self.assertEqual(added["data"]["added"], 1)
        again = self.client.post(
            f"/api/courses/{course_id}/students",
            headers=teacher_headers,
            json={"student_ids": [student_id]},
        ).json()
        self.assertEqual(again["data"]["added"], 0)
        roster = self.client.get(f"/api/courses/{course_id}/students", headers=teacher_headers).json()["data"]
        self.assertEqual([item["student_id"] for item in roster], [student_id])

        # 7) 学生只能看课表，不能建课/管理名单；「我的课程」按选课过滤
        forbidden = self.client.post(
            "/api/courses",
            headers=student_headers,
            json={
                "name": "越权课程",
                "lab_id": lab_id,
                "weekday": weekday,
                "start_time": "14:00:00",
                "end_time": "15:00:00",
                "start_date": today.isoformat(),
                "end_date": (today + timedelta(days=60)).isoformat(),
            },
        )
        self.assertEqual(forbidden.status_code, 403)
        self.assertEqual(
            self.client.post(
                f"/api/courses/{course_id}/students",
                headers=student_headers,
                json={"student_ids": [student_id]},
            ).status_code,
            403,
        )
        mine = self.client.get("/api/courses", headers=student_headers, params={"mine": True}).json()["data"]
        self.assertEqual([item["id"] for item in mine], [course_id])

        # 8) 移出名单
        self.assertEqual(
            self.client.delete(f"/api/courses/{course_id}/students/{student_id}", headers=teacher_headers).status_code,
            200,
        )

    def test_equipment_usage_doc_url_roundtrip(self):
        """设备操作说明：创建/列表/更新往返，站内路径与清空可用，危险协议被拒绝。"""
        admin_headers = self.login("admin", "Admin123456")

        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "说明书实验室", "location": "实训楼 502", "capacity": 10},
        ).json()
        self.assertEqual(lab["code"], 200)
        lab_id = lab["data"]["id"]

        created = self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={
                "lab_id": lab_id,
                "name": "示波器",
                "code": "DOC-OSC-001",
                "quantity": 2,
                "description": "双通道数字示波器",
                "usage_doc_url": "https://example.com/docs/oscilloscope.pdf",
            },
        ).json()
        self.assertEqual(created["code"], 200)
        equipment_id = created["data"]["id"]
        self.assertEqual(created["data"]["usage_doc_url"], "https://example.com/docs/oscilloscope.pdf")

        # 列表接口一并返回操作说明（学生端也靠它展示「查看说明」）
        listed = self.client.get("/api/equipment", headers=admin_headers, params={"lab_id": lab_id}).json()
        row = next(item for item in listed["data"] if item["id"] == equipment_id)
        self.assertEqual(row["usage_doc_url"], "https://example.com/docs/oscilloscope.pdf")

        # 站内绝对路径（如上传目录里的 PDF）同样允许
        updated = self.client.put(
            f"/api/equipment/{equipment_id}",
            headers=admin_headers,
            json={"usage_doc_url": " /uploads/manual.pdf "},
        ).json()
        self.assertEqual(updated["code"], 200)
        self.assertEqual(updated["data"]["usage_doc_url"], "/uploads/manual.pdf")

        # 空字符串视为清空
        cleared = self.client.put(
            f"/api/equipment/{equipment_id}",
            headers=admin_headers,
            json={"usage_doc_url": ""},
        ).json()
        self.assertEqual(cleared["code"], 200)
        self.assertIsNone(cleared["data"]["usage_doc_url"])

        # 危险协议（javascript: 等）必须拒绝，避免前端渲染成可点击链接
        bad = self.client.put(
            f"/api/equipment/{equipment_id}",
            headers=admin_headers,
            json={"usage_doc_url": "javascript:alert(1)"},
        )
        self.assertEqual(bad.status_code, 400)

    def test_password_change_invalidates_old_token(self):
        """修改密码后旧 JWT 必须失效，新密码可以重新登录。"""
        register_data = {
            "username": "password_user",
            "password": "OldPassword123",
            "email": "password-user@example.com",
            "full_name": "改密测试用户",
        }
        self.assertEqual(
            self.client.post("/api/auth/register", json=register_data).status_code,
            200,
        )
        old_headers = self.login("password_user", "OldPassword123")

        changed = self.client.put(
            "/api/user/password",
            headers=old_headers,
            json={"old_password": "OldPassword123", "new_password": "NewPassword456"},
        )
        self.assertEqual(changed.status_code, 200)

        old_token_result = self.client.get("/api/user/info", headers=old_headers)
        self.assertEqual(old_token_result.status_code, 401)
        self.assertEqual(old_token_result.json()["code"], 401)

        new_headers = self.login("password_user", "NewPassword456")
        self.assertEqual(
            self.client.get("/api/user/info", headers=new_headers).status_code,
            200,
        )

    def test_ai_confirmation_uses_latest_request_and_respects_cancellation(self):
        """多轮确认应选最新参数；否定或撤回消息绝不能调用写工具。"""
        old_request = "预约实验室ID 1 2026-10-01 09:00-10:00 用途：旧需求"
        new_request = "预约实验室ID 2 2026-10-02 11:00-12:00 用途：新需求"
        history = [
            {"role": "user", "content": old_request},
            {"role": "user", "content": new_request},
        ]
        self.assertFalse(has_explicit_reservation_intent("不要帮我预约"))
        self.assertFalse(has_explicit_reservation_intent("我不想创建预约"))

        with patch("app.ai.agent.create_lab_reservation", return_value="预约已创建") as create_tool:
            local_tool_node({"message": "确认预约", "history": history, "user_id": 1, "db": object()})
            self.assertEqual(create_tool.call_args.args[2]["lab_id"], 2)
            self.assertEqual(create_tool.call_args.args[2]["purpose"], "新需求")

            create_tool.reset_mock()
            local_tool_node({"message": "不要帮我预约", "history": history, "user_id": 1, "db": object()})
            create_tool.assert_not_called()

            cancelled_history = [*history, {"role": "user", "content": "取消预约"}]
            local_tool_node({"message": "确认预约", "history": cancelled_history, "user_id": 1, "db": object()})
            create_tool.assert_not_called()

            stale_history = [history[0], {"role": "user", "content": "查询设备故障说明"}]
            local_tool_node({"message": "确认预约", "history": stale_history, "user_id": 1, "db": object()})
            create_tool.assert_not_called()

    def test_full_business_flow(self):
        # 1. 注册、重复注册与登录。
        register_data = {
            "username": "student01",
            "password": "Student123",
            "email": "student01@example.com",
            "full_name": "测试学生",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 400)

        # 错误密码必须返回真实 HTTP 401，同时仍保留统一响应体。
        invalid_login = self.client.post(
            "/api/auth/login",
            json={"username": "student01", "password": "WrongPassword123"},
        )
        self.assertEqual(invalid_login.status_code, 401)
        self.assertEqual(invalid_login.json()["code"], 401)

        student_headers = self.login("student01", "Student123")
        admin_headers = self.login("admin", "Admin123456")

        invalid_user = self.client.post(
            "/api/user",
            headers=admin_headers,
            json={"username": "   ", "password": "Password123", "email": "invalid-name@example.com"},
        )
        self.assertEqual(invalid_user.status_code, 400)

        info = self.client.get("/api/user/info", headers=student_headers).json()
        self.assertEqual(info["data"]["full_name"], "测试学生")

        # 普通用户不能调用管理员接口；依赖层应返回真实 HTTP 403。
        forbidden = self.client.post(
            "/api/labs",
            headers=student_headers,
            json={"name": "越权实验室", "location": "未知位置", "capacity": 1},
        )
        self.assertEqual(forbidden.status_code, 403)
        self.assertEqual(forbidden.json()["code"], 403)

        # 2. 管理员创建实验室和设备。
        lab_response = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={
                "name": "  人工智能实验室  ",
                "location": "实训楼 301",
                "capacity": 40,
                "description": "AI 实验教学",
                "status": 1,
            },
        ).json()
        self.assertEqual(lab_response["code"], 200)
        self.assertEqual(lab_response["data"]["name"], "人工智能实验室")
        lab_id = lab_response["data"]["id"]

        # 显式 null 和全空格不能绕过更新/新增校验写坏数据库必填列。
        self.assertEqual(
            self.client.put(f"/api/labs/{lab_id}", headers=admin_headers, json={"status": None}).status_code,
            400,
        )
        self.assertEqual(
            self.client.post("/api/labs", headers=admin_headers, json={"name": "  ", "location": "三楼"}).status_code,
            400,
        )

        equipment_response = self.client.post(
            "/api/equipment",
            headers=admin_headers,
            json={
                "lab_id": lab_id,
                "name": " GPU 工作站 ",
                "code": " GPU-001 ",
                "quantity": 10,
                "description": "模型训练",
                "status": 1,
            },
        ).json()
        self.assertEqual(equipment_response["code"], 200)
        self.assertEqual(equipment_response["data"]["code"], "GPU-001")
        self.assertEqual(
            self.client.put(
                f"/api/equipment/{equipment_response['data']['id']}",
                headers=admin_headers,
                json={"quantity": None},
            ).status_code,
            400,
        )

        # 3. 学生提交预约；同一时间段的第二条申请会被拦截。
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        reservation_data = {
            "lab_id": lab_id,
            "reservation_date": tomorrow,
            "start_time": "09:00:00",
            "end_time": "11:00:00",
            "purpose": "人工智能实验教学",
        }
        reservation_response = self.client.post(
            "/api/reservations", headers=student_headers, json=reservation_data
        ).json()
        self.assertEqual(reservation_response["code"], 200)
        # 创建接口支持每周重复，返回创建成功的预约列表；此处未开启重复，只有一条。
        self.assertEqual(len(reservation_response["data"]), 1)
        reservation_id = reservation_response["data"][0]["id"]
        conflict = self.client.post("/api/reservations", headers=student_headers, json=reservation_data).json()
        self.assertEqual(conflict["code"], 400)

        past_time_today = {
            **reservation_data,
            "reservation_date": date.today().isoformat(),
            "start_time": "00:00:00",
            "end_time": "00:01:00",
        }
        self.assertEqual(
            self.client.post("/api/reservations", headers=student_headers, json=past_time_today).status_code,
            400,
        )

        # 4. 实验室暂停后不能批准；恢复开放后，管理员可以正常审核。
        self.assertEqual(
            self.client.put(f"/api/labs/{lab_id}", headers=admin_headers, json={"status": 0}).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                f"/api/reservations/{reservation_id}/review",
                headers=admin_headers,
                json={"status": "approved", "review_comment": "暂停期间不应通过"},
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.put(f"/api/labs/{lab_id}", headers=admin_headers, json={"status": 1}).status_code,
            200,
        )

        # 5. 管理员审核后，学生可以在“我的预约”看到结果并取消。
        review = self.client.put(
            f"/api/reservations/{reservation_id}/review",
            headers=admin_headers,
            json={"status": "approved", "review_comment": "请按时到场"},
        ).json()
        self.assertEqual(review["code"], 200)
        mine = self.client.get("/api/reservations/mine", headers=student_headers).json()
        self.assertEqual(mine["data"][0]["status"], "approved")
        cancel = self.client.put(f"/api/reservations/{reservation_id}/cancel", headers=student_headers).json()
        self.assertEqual(cancel["code"], 200)

        # 5. 上传接口与本地 AI SSE 链路都能独立运行。
        invalid_upload = self.client.post(
            "/api/upload/image",
            headers=student_headers,
            files={"file": ("demo.png", b"fake-png-content", "image/png")},
        )
        self.assertEqual(invalid_upload.status_code, 400)

        # 使用真实的 1×1 PNG，验证合法图片仍能保存并返回静态路径。
        valid_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory(prefix="lab_agent_upload_test_") as upload_dir:
            with patch("app.api.upload.UPLOAD_DIR", Path(upload_dir)):
                upload = self.client.post(
                    "/api/upload/image",
                    headers=student_headers,
                    files={"file": ("demo.png", valid_png, "image/png")},
                ).json()
                self.assertEqual(upload["code"], 200)
                self.assertTrue((Path(upload_dir) / upload["data"]["filename"]).exists())

        ai_response = self.client.post(
            "/api/ai/chat/stream",
            headers=student_headers,
            json={"message": "当前有哪些开放的实验室和设备？", "history": []},
        )
        self.assertEqual(ai_response.status_code, 200)
        self.assertIn("event: token", ai_response.text)
        # SSE 会把回答拆成多个 token 事件，拼接完整回答后再校验内容。
        streamed_answer = extract_streamed_tokens(ai_response.text)
        self.assertIn("GPU", streamed_answer)

        # 6. 多轮 AI 预约：上一轮提供参数，本轮明确确认；重复确认不能生成第二条记录。
        reservation_instruction = f"预约实验室ID {lab_id} {tomorrow} 13:00-14:00 用途：多轮对话预约测试"
        multi_turn_payload = {
            "message": "确认预约",
            "history": [{"role": "user", "content": reservation_instruction}],
        }
        ai_reservation = self.client.post(
            "/api/ai/chat/stream",
            headers=student_headers,
            json=multi_turn_payload,
        )
        self.assertEqual(ai_reservation.status_code, 200)
        self.assertIn("create_lab_reservation", ai_reservation.text)

        # 再发一次相同确认会命中时间冲突，数据库中仍然只有一条该用途的预约。
        self.client.post(
            "/api/ai/chat/stream",
            headers=student_headers,
            json=multi_turn_payload,
        )
        mine_after_ai = self.client.get(
            "/api/reservations/mine",
            headers=student_headers,
        ).json()["data"]
        ai_reservations = [item for item in mine_after_ai if item["purpose"] == "多轮对话预约测试"]
        self.assertEqual(len(ai_reservations), 1)
        self.assertEqual(ai_reservations[0]["status"], "pending")

    def test_avatar_persists_and_plain_profile_save_keeps_it(self):
        """头像上传后必须落库；之后再保存普通资料不能把头像清空。"""
        register_data = {
            "username": "avatar_user",
            "password": "Avatar12345",
            "email": "avatar-user@example.com",
            "full_name": "头像用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("avatar_user", "Avatar12345")

        valid_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory(prefix="lab_agent_avatar_test_") as upload_dir:
            with patch("app.api.upload.UPLOAD_DIR", Path(upload_dir)):
                upload = self.client.post(
                    "/api/upload/image",
                    headers=headers,
                    files={"file": ("avatar.png", valid_png, "image/png")},
                ).json()
                self.assertEqual(upload["code"], 200)
        avatar_url = f"http://127.0.0.1:8000{upload['data']['url']}"

        saved = self.client.put(
            "/api/user/profile",
            headers=headers,
            json={"email": "avatar-user@example.com", "full_name": "头像用户", "avatar_url": avatar_url},
        ).json()
        self.assertEqual(saved["code"], 200)
        self.assertEqual(saved["data"]["avatar_url"], avatar_url)

        plain_save = self.client.put(
            "/api/user/profile",
            headers=headers,
            json={"email": "avatar-user@example.com", "full_name": "改个名字"},
        ).json()
        self.assertEqual(plain_save["code"], 200)
        # 没有显式携带 avatar_url 时，已有头像不能被清空。
        self.assertEqual(plain_save["data"]["avatar_url"], avatar_url)

    def test_reservation_update_cannot_bypass_slot_rules(self):
        """修改待审核预约时，冲突/当天过期/超时长/黑名单规则必须生效。"""
        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "editor_user",
            "password": "Editor12345",
            "email": "editor-user@example.com",
            "full_name": "修改预约用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("editor_user", "Editor12345")

        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "修改校验实验室", "location": "实训楼 501", "capacity": 10, "status": 1},
        ).json()["data"]["id"]

        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        first = self.client.post(
            "/api/reservations",
            headers=headers,
            json={
                "lab_id": lab_id,
                "reservation_date": tomorrow,
                "start_time": "09:00:00",
                "end_time": "11:00:00",
                "purpose": "修改校验-第一条",
            },
        ).json()
        self.assertEqual(first["code"], 200)
        second_id = self.client.post(
            "/api/reservations",
            headers=headers,
            json={
                "lab_id": lab_id,
                "reservation_date": tomorrow,
                "start_time": "11:00:00",
                "end_time": "12:00:00",
                "purpose": "修改校验-第二条",
            },
        ).json()["data"][0]["id"]

        # 1) 改成与已有预约冲突的时间
        conflict = self.client.put(
            f"/api/reservations/{second_id}",
            headers=headers,
            json={"start_time": "10:00:00", "end_time": "11:30:00"},
        )
        self.assertEqual(conflict.status_code, 400)

        # 2) 改成今天已经过去的开始时间
        past = self.client.put(
            f"/api/reservations/{second_id}",
            headers=headers,
            json={"reservation_date": date.today().isoformat(), "start_time": "00:00:00", "end_time": "00:30:00"},
        )
        self.assertEqual(past.status_code, 400)

        # 3) 单次超过 4 小时
        too_long = self.client.put(
            f"/api/reservations/{second_id}",
            headers=headers,
            json={"start_time": "12:00:00", "end_time": "17:00:00"},
        )
        self.assertEqual(too_long.status_code, 400)

        # 4) 改成管理员设置的不可约时段
        blackout = self.client.post(
            "/api/blackouts",
            headers=admin_headers,
            json={
                "lab_id": lab_id,
                "blackout_date": tomorrow,
                "start_time": "13:00:00",
                "end_time": "14:00:00",
                "reason": "维护测试",
            },
        ).json()
        self.assertEqual(blackout["code"], 200)
        blocked = self.client.put(
            f"/api/reservations/{second_id}",
            headers=headers,
            json={"start_time": "13:00:00", "end_time": "13:30:00"},
        )
        self.assertEqual(blocked.status_code, 400)

        # 所有非法修改都被拦截后，原记录保持 11:00-12:00 不变
        mine = self.client.get("/api/reservations/mine", headers=headers).json()["data"]
        current = next(item for item in mine if item["id"] == second_id)
        self.assertEqual(str(current["start_time"]), "11:00:00")
        self.assertEqual(str(current["end_time"]), "12:00:00")

    def test_reservation_calendar_returns_range_events_and_blackouts(self):
        """周历接口：日期范围缺省取当前周一、事件与不可约时段按范围返回。"""
        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "calendar_user",
            "password": "Calendar12345",
            "email": "calendar-user@example.com",
            "full_name": "周历测试用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("calendar_user", "Calendar12345")

        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "周历测试实验室", "location": "实训楼 701", "capacity": 12, "status": 1},
        ).json()["data"]["id"]

        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        self.assertEqual(
            self.client.post(
                "/api/reservations",
                headers=headers,
                json={
                    "lab_id": lab_id,
                    "reservation_date": tomorrow,
                    "start_time": "09:00:00",
                    "end_time": "11:00:00",
                    "purpose": "周历占用展示",
                },
            ).json()["code"],
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/blackouts",
                headers=admin_headers,
                json={
                    "lab_id": lab_id,
                    "blackout_date": tomorrow,
                    "start_time": "13:00:00",
                    "end_time": "14:00:00",
                    "reason": "周历维护",
                },
            ).json()["code"],
            200,
        )

        # 显式单日范围：事件与不可约时段都要返回
        day_view = self.client.get(
            "/api/reservations/calendar",
            headers=headers,
            params={"lab_id": lab_id, "start_date": tomorrow, "end_date": tomorrow},
        ).json()
        self.assertEqual(day_view["code"], 200)
        self.assertEqual(day_view["data"]["range"], {"start_date": tomorrow, "end_date": tomorrow})
        self.assertEqual(day_view["data"]["lab"]["name"], "周历测试实验室")
        self.assertEqual(len(day_view["data"]["reservations"]), 1)
        event = day_view["data"]["reservations"][0]
        self.assertEqual(event["reservation_date"], tomorrow)
        self.assertEqual(str(event["start_time"]), "09:00:00")
        self.assertEqual(str(event["end_time"]), "11:00:00")
        self.assertEqual(event["purpose"], "周历占用展示")
        self.assertEqual(len(day_view["data"]["blackouts"]), 1)
        self.assertEqual(day_view["data"]["blackouts"][0]["reason"], "周历维护")

        # 缺省范围：从本周周一开始，共 7 天
        default_view = self.client.get("/api/reservations/calendar", headers=headers, params={"lab_id": lab_id}).json()
        self.assertEqual(default_view["code"], 200)
        start = date.fromisoformat(default_view["data"]["range"]["start_date"])
        end = date.fromisoformat(default_view["data"]["range"]["end_date"])
        self.assertEqual(start.weekday(), 0)  # 一周从周一开始
        self.assertEqual((end - start).days, 6)

        # 非法范围与超长范围都被拦截；实验室不存在返回真实 404
        reversed_range = self.client.get(
            "/api/reservations/calendar",
            headers=headers,
            params={"lab_id": lab_id, "start_date": tomorrow, "end_date": date.today().isoformat()},
        )
        self.assertEqual(reversed_range.status_code, 400)
        too_long = self.client.get(
            "/api/reservations/calendar",
            headers=headers,
            params={
                "lab_id": lab_id,
                "start_date": tomorrow,
                "end_date": (date.today() + timedelta(days=40)).isoformat(),
            },
        )
        self.assertEqual(too_long.status_code, 400)
        missing_lab = self.client.get("/api/reservations/calendar", headers=headers, params={"lab_id": 999999})
        self.assertEqual(missing_lab.status_code, 404)

        # 已取消的预约不再占用日历时间线
        reservation_id = self.client.get(
            "/api/reservations/mine", headers=headers, params={"status": "pending"}
        ).json()["data"][0]["id"]
        self.assertEqual(
            self.client.put(f"/api/reservations/{reservation_id}/cancel", headers=headers).json()["code"],
            200,
        )
        after_cancel = self.client.get(
            "/api/reservations/calendar",
            headers=headers,
            params={"lab_id": lab_id, "start_date": tomorrow, "end_date": tomorrow},
        ).json()
        self.assertEqual(after_cancel["data"]["reservations"], [])
        self.assertEqual(len(after_cancel["data"]["blackouts"]), 1)

    def test_stats_dashboard_returns_trend_usage_and_repairs(self):
        """仪表盘接口：趋势按天补零、实验室排行、报修分布，且学生只能看到自己的数据。"""
        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "dashboard_user",
            "password": "Dashboard12345",
            "email": "dashboard-user@example.com",
            "full_name": "仪表盘测试用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("dashboard_user", "Dashboard12345")

        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "统计图表实验室", "location": "实训楼 801", "capacity": 6, "status": 1},
        ).json()["data"]["id"]

        # 学生创建 1 条预约（create_time 为今天，应落在趋势图的今天桶里）
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        self.assertEqual(
            self.client.post(
                "/api/reservations",
                headers=headers,
                json={
                    "lab_id": lab_id,
                    "reservation_date": tomorrow,
                    "start_time": "09:00:00",
                    "end_time": "10:00:00",
                    "purpose": "仪表盘趋势展示",
                },
            ).json()["code"],
            200,
        )
        # 学生创建 2 条报修，管理员处理掉 1 条
        repair_ids = [
            self.client.post(
                "/api/repairs",
                headers=headers,
                json={"title": f"仪表盘报修 {i}", "description": "用于统计图表验证。"},
            ).json()["data"]["id"]
            for i in (1, 2)
        ]
        self.assertEqual(
            self.client.put(
                f"/api/repairs/{repair_ids[0]}/handle",
                headers=admin_headers,
                json={"status": "resolved", "handle_comment": "已修复"},
            ).json()["code"],
            200,
        )

        # 管理员视角：默认 7 天，最后一天是今天
        admin_view = self.client.get("/api/stats/dashboard", headers=admin_headers).json()
        self.assertEqual(admin_view["code"], 200)
        trend = admin_view["data"]["reservation_trend"]
        self.assertEqual(len(trend), 7)
        today_entry = trend[-1]
        self.assertEqual(today_entry["date"], date.today().isoformat())
        # 测试类共享一个数据库，其他用例也会产生预约，这里只断言“今天至少有 1 条”
        self.assertGreaterEqual(
            today_entry["pending"] + today_entry["approved"] + today_entry["rejected"] + today_entry["cancelled"],
            1,
        )
        usage_names = [item["lab_name"] for item in admin_view["data"]["lab_usage"]]
        self.assertIn("统计图表实验室", usage_names)
        repair_map = {item["status"]: item["count"] for item in admin_view["data"]["repair_status"]}
        self.assertGreaterEqual(repair_map.get("pending", 0), 1)
        self.assertGreaterEqual(repair_map.get("resolved", 0), 1)

        # 学生视角：报修分布只统计自己（1 条待处理 + 1 条已解决）
        student_view = self.client.get("/api/stats/dashboard", headers=headers).json()
        self.assertEqual(student_view["code"], 200)
        student_repair_total = sum(item["count"] for item in student_view["data"]["repair_status"])
        self.assertEqual(student_repair_total, 2)
        student_usage_names = [item["lab_name"] for item in student_view["data"]["lab_usage"]]
        self.assertIn("统计图表实验室", student_usage_names)

        # days 参数：30 天返回 30 个点，超界收敛到 90
        self.assertEqual(
            len(
                self.client.get("/api/stats/dashboard", headers=admin_headers, params={"days": 30}).json()["data"][
                    "reservation_trend"
                ]
            ),
            30,
        )
        self.assertEqual(
            len(
                self.client.get("/api/stats/dashboard", headers=admin_headers, params={"days": 500}).json()["data"][
                    "reservation_trend"
                ]
            ),
            90,
        )

    def test_scheduler_expires_stale_and_reminds_upcoming(self):
        """定时任务纯函数：过期预约自动关闭并通知；开课前提醒只发一次。"""
        from app.services.schedule import expire_stale_reservations, remind_upcoming_reservations

        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "schedule_user",
            "password": "Schedule12345",
            "email": "schedule-user@example.com",
            "full_name": "定时任务用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("schedule_user", "Schedule12345")

        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "定时任务实验室", "location": "实训楼 901", "capacity": 6, "status": 1},
        ).json()["data"]["id"]

        # —— 过期场景：接口层不允许创建过去的预约，这里直写数据库构造“昨天”的记录 ——
        # 注意：commit 会让 ORM 实例属性过期，会话关闭后再访问属性会抛
        # DetachedInstanceError，所以 user_id 必须在会话内先读出来存成普通变量。
        with SessionLocal() as db:
            user_id = db.query(User.id).filter(User.username == "schedule_user").scalar()
            stale = Reservation(
                user_id=user_id,
                lab_id=lab_id,
                reservation_date=date.today() - timedelta(days=1),
                start_time=time(9, 0),
                end_time=time(10, 0),
                purpose="验证过期自动关闭",
                status="approved",
            )
            db.add(stale)
            db.commit()
            stale_id = stale.id

        with SessionLocal() as db:
            affected = expire_stale_reservations(db)
        self.assertGreaterEqual(affected, 1)
        with SessionLocal() as db:
            self.assertEqual(db.get(Reservation, stale_id).status, "expired")
        notifications = self.client.get("/api/notifications/mine", headers=headers).json()["data"]
        expired_notice = next(
            (n for n in notifications if n["related_id"] == stale_id and n["type"] == "reservation_result"),
            None,
        )
        self.assertIsNotNone(expired_notice)
        self.assertIn("自动关闭", expired_notice["title"])

        # —— 提醒场景：今天 40 分钟后开始的已通过预约，应收到一次提醒 ——
        starts_at = datetime.now() + timedelta(minutes=40)
        with SessionLocal() as db:
            upcoming = Reservation(
                user_id=user_id,
                lab_id=lab_id,
                reservation_date=starts_at.date(),
                start_time=starts_at.time().replace(microsecond=0),
                end_time=(starts_at + timedelta(minutes=60)).time().replace(microsecond=0),
                purpose="验证开课前提醒",
                status="approved",
            )
            db.add(upcoming)
            db.commit()
            upcoming_id = upcoming.id

        with SessionLocal() as db:
            self.assertEqual(remind_upcoming_reservations(db), 1)
        notifications = self.client.get("/api/notifications/mine", headers=headers).json()["data"]
        reminders = [n for n in notifications if n["type"] == "reservation_reminder" and n["related_id"] == upcoming_id]
        self.assertEqual(len(reminders), 1)
        self.assertIn("即将开始", reminders[0]["title"])

        # 幂等：再次执行不重复发提醒
        with SessionLocal() as db:
            self.assertEqual(remind_upcoming_reservations(db), 0)

    def test_websocket_pushes_review_result_and_rejects_bad_token(self):
        """WebSocket：审核结果实时推给在线用户；非法 token 以 4401 关闭。"""
        from starlette.websockets import WebSocketDisconnect

        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "ws_user",
            "password": "WsUser123456",
            "email": "ws-user@example.com",
            "full_name": "实时通知用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        login_res = self.client.post("/api/auth/login", json={"username": "ws_user", "password": "WsUser123456"}).json()
        token = login_res["data"]["token"]

        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "实时通知实验室", "location": "实训楼 1001", "capacity": 6, "status": 1},
        ).json()["data"]["id"]
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        reservation_id = self.client.post(
            "/api/reservations",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "lab_id": lab_id,
                "reservation_date": tomorrow,
                "start_time": "10:00:00",
                "end_time": "11:00:00",
                "purpose": "验证 WebSocket 推送",
            },
        ).json()["data"][0]["id"]

        # 非法 token：服务端 accept 后立即以 4401 关闭，读取消息时抛 WebSocketDisconnect
        with self.client.websocket_connect("/api/ws/notifications?token=not-a-jwt") as ws:
            with self.assertRaises(WebSocketDisconnect):
                ws.receive_json()

        # 合法 token：连接保持期间，管理员审核通过的结果应实时推送过来
        with self.client.websocket_connect(f"/api/ws/notifications?token={token}") as ws:
            review = self.client.put(
                f"/api/reservations/{reservation_id}/review",
                headers=admin_headers,
                json={"status": "approved", "review_comment": "同意"},
            )
            self.assertEqual(review.json()["code"], 200)
            message = ws.receive_json()
        self.assertEqual(message["kind"], "notification")
        self.assertEqual(message["notification"]["type"], "reservation_result")
        self.assertIn("已通过", message["notification"]["title"])
        self.assertEqual(message["notification"]["related_id"], reservation_id)

    def test_login_captcha_lockout_and_audit_log(self):
        """登录安全：验证码一次性校验、失败锁定、登录日志落库与管理员查询。"""
        from app.config import settings
        from app.models.login_log import LoginLog
        from app.services import login_security
        from app.services.login_security import _CAPTCHAS

        username = "sec_user"
        res = self.client.post(
            "/api/auth/register",
            json={
                "username": username,
                "password": "SecUser123456",
                "email": "sec_user@example.com",
                "full_name": "安全测试",
            },
        )
        self.assertEqual(res.json()["code"], 200)

        original = settings.LOGIN_CAPTCHA_ENABLED
        try:
            # 1) 开启验证码后：不带验证码直接拒绝
            settings.LOGIN_CAPTCHA_ENABLED = True
            res = self.client.post("/api/auth/login", json={"username": username, "password": "SecUser123456"})
            self.assertEqual(res.json()["code"], 400)
            self.assertIn("验证码", res.json()["message"])

            # 2) 错误验证码 → 400（captcha_id 不存在）
            res = self.client.post(
                "/api/auth/login",
                json={
                    "username": username,
                    "password": "SecUser123456",
                    "captcha_id": "not-exist",
                    "captcha_code": "AAAA",
                },
            )
            self.assertEqual(res.json()["code"], 400)

            # 3) 取真实验证码（测试直接从内存读答案，不做 OCR），登录成功
            captcha_res = self.client.get("/api/auth/captcha").json()
            self.assertTrue(captcha_res["data"]["enabled"])
            code = _CAPTCHAS[captcha_res["data"]["captcha_id"]][0]
            res = self.client.post(
                "/api/auth/login",
                json={
                    "username": username,
                    "password": "SecUser123456",
                    "captcha_id": captcha_res["data"]["captcha_id"],
                    "captcha_code": code,
                },
            )
            self.assertEqual(res.json()["code"], 200)

            # 4) 验证码一次性消费：二次使用同一 captcha_id 失败
            res = self.client.post(
                "/api/auth/login",
                json={
                    "username": username,
                    "password": "SecUser123456",
                    "captcha_id": captcha_res["data"]["captcha_id"],
                    "captcha_code": code,
                },
            )
            self.assertEqual(res.json()["code"], 400)
        finally:
            settings.LOGIN_CAPTCHA_ENABLED = original

        # 5) 失败锁定：连续 5 次错误后，正确密码也拒绝（403 + 锁定提示）
        for _ in range(settings.LOGIN_MAX_FAILURES):
            res = self.client.post("/api/auth/login", json={"username": username, "password": "WrongPassword123"})
            self.assertEqual(res.json()["code"], 401)
        res = self.client.post("/api/auth/login", json={"username": username, "password": "SecUser123456"})
        self.assertEqual(res.json()["code"], 403)
        self.assertIn("锁定", res.json()["message"])

        # 6) 成功登录会清空失败计数：解锁后可正常登录
        login_security.clear_failures(username)
        res = self.client.post("/api/auth/login", json={"username": username, "password": "SecUser123456"})
        self.assertEqual(res.json()["code"], 200)

        # 7) 登录日志：失败与成功都有落库记录（含 IP / UA）
        with SessionLocal() as db:
            logs = db.query(LoginLog).filter(LoginLog.username == username).all()
            self.assertTrue(any(log.success for log in logs))
            self.assertTrue(any((not log.success) and "账号或密码错误" in log.message for log in logs))

            # 8) 管理员可查询审计日志并按用户名过滤
            admin_login = self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "Admin123456"}
            ).json()
            headers = {"Authorization": f"Bearer {admin_login['data']['token']}"}
            res = self.client.get("/api/auth/login-logs", headers=headers, params={"username": username})
            self.assertEqual(res.json()["code"], 200)
            payload = res.json()["data"]
            self.assertGreaterEqual(payload["total"], 1)
            self.assertTrue(all(item["username"] == username for item in payload["items"]))

            # 9) 普通用户无权查看审计日志
            user_login = self.client.post(
                "/api/auth/login", json={"username": username, "password": "SecUser123456"}
            ).json()
            res = self.client.get(
                "/api/auth/login-logs", headers={"Authorization": f"Bearer {user_login['data']['token']}"}
            )
            self.assertEqual(res.json()["code"], 403)

    def test_repair_handling_creates_notification(self):
        """管理员处理报修后，报修人应收到站内通知。"""
        admin_headers = self.login("admin", "Admin123456")
        register_data = {
            "username": "repair_user",
            "password": "Repair12345",
            "email": "repair-user@example.com",
            "full_name": "报修用户",
        }
        self.assertEqual(self.client.post("/api/auth/register", json=register_data).json()["code"], 200)
        headers = self.login("repair_user", "Repair12345")

        created = self.client.post(
            "/api/repairs",
            headers=headers,
            json={"title": "投影仪无法开机", "description": "按电源键没有反应，指示灯不亮。"},
        ).json()
        self.assertEqual(created["code"], 200)
        repair_id = created["data"]["id"]

        handled = self.client.put(
            f"/api/repairs/{repair_id}/handle",
            headers=admin_headers,
            json={"status": "processing", "handle_comment": "已联系维修师傅，明天上门。"},
        ).json()
        self.assertEqual(handled["code"], 200)

        notifications = self.client.get("/api/notifications/mine", headers=headers).json()["data"]
        repair_notifications = [item for item in notifications if item["type"] == "repair_result"]
        self.assertEqual(len(repair_notifications), 1)
        self.assertIn("处理中", repair_notifications[0]["title"])
        self.assertEqual(repair_notifications[0]["related_id"], repair_id)

    def test_ai_recommendation_respects_blackout_and_past_dates(self):
        """推荐工具应跳过不可约时段和已经过去的日期。"""
        from app.ai.tools import check_lab_availability, recommend_lab

        admin_headers = self.login("admin", "Admin123456")
        lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "黑名单推荐实验室", "location": "实训楼 601", "capacity": 8, "status": 1},
        ).json()["data"]["id"]
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        self.client.post(
            "/api/blackouts",
            headers=admin_headers,
            json={
                "lab_id": lab_id,
                "blackout_date": tomorrow,
                "start_time": "13:00:00",
                "end_time": "18:00:00",
                "reason": "全天维护",
            },
        )

        with SessionLocal() as db:
            availability = check_lab_availability(db, lab_id, tomorrow)
            self.assertIn("不可约", availability)
            self.assertIn("全天维护", availability)

            # 该实验室整个下午都不可约，不能被推荐
            recommendation = recommend_lab(db, tomorrow, "afternoon", 2)
            self.assertNotIn("黑名单推荐实验室", recommendation)

            # 已经过去的日期不再给出推荐
            yesterday = (date.today() - timedelta(days=1)).isoformat()
            past = recommend_lab(db, yesterday, "afternoon", 2)
            self.assertIn("已经过去", past)

    def test_approval_flow_multi_step_and_reset_on_edit(self):
        """多级审批：按节点流转、当前节点权限、驳回、改约重走与流程配置约束。"""
        from app.models.notification import Notification
        from app.models.reservation import Reservation

        admin_headers = self.login("admin", "Admin123456")
        lab = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "多级审批测试实验室", "location": "实训楼 901", "capacity": 10},
        ).json()["data"]
        lab_id = lab["id"]

        def create_teacher(username: str) -> int:
            response = self.client.post(
                "/api/user",
                headers=admin_headers,
                json={
                    "username": username,
                    "password": "Teacher123456",
                    "email": f"{username}@example.com",
                    "is_teacher": True,
                },
            ).json()
            self.assertEqual(response["code"], 200)
            return response["data"]["id"]

        teacher_a_id = create_teacher("flow_teacher_a")
        create_teacher("flow_teacher_b")
        teacher_a_headers = self.login("flow_teacher_a", "Teacher123456")
        teacher_b_headers = self.login("flow_teacher_b", "Teacher123456")

        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={
                    "username": "flow_student",
                    "password": "FlowStu123456",
                    "email": "flow_student@example.com",
                },
            ).json()["code"],
            200,
        )
        student_headers = self.login("flow_student", "FlowStu123456")

        flow_payload = {
            "name": "重点实验室两级审批",
            "description": "导师初审后由管理员终审",
            "enabled": True,
            "steps": [
                {
                    "name": "导师初审",
                    "approver_type": "user",
                    "approver_id": teacher_a_id,
                },
                {
                    "name": "管理员终审",
                    "approver_type": "role",
                    "approver_role": "admin",
                },
            ],
        }
        saved = self.client.put(f"/api/approval-flows/{lab_id}", headers=admin_headers, json=flow_payload).json()
        self.assertEqual(saved["code"], 200)
        self.assertEqual(saved["data"]["lab_name"], "多级审批测试实验室")
        self.assertEqual([step["label"] for step in saved["data"]["steps"]], ["导师初审", "管理员终审"])
        flow_id = saved["data"]["id"]

        # 审批流配置是管理员能力；普通用户不能查看/修改
        self.assertEqual(self.client.get("/api/approval-flows", headers=student_headers).status_code, 403)
        self.assertEqual(
            self.client.put(f"/api/approval-flows/{lab_id}", headers=teacher_a_headers, json=flow_payload).status_code,
            403,
        )
        invalid_steps = {**flow_payload, "steps": []}
        self.assertEqual(
            self.client.put(f"/api/approval-flows/{lab_id}", headers=admin_headers, json=invalid_steps).status_code,
            400,
        )

        today = date.today()
        first_date = today + timedelta(days=7)
        created = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": first_date.isoformat(),
                "start_time": "09:00:00",
                "end_time": "10:00:00",
                "purpose": "多级审批第一单",
            },
        ).json()
        self.assertEqual(created["code"], 200)
        reservation_id = created["data"][0]["id"]
        approval = created["data"][0]["approval"]
        self.assertEqual(approval["current_step"], 1)
        self.assertEqual(approval["total_steps"], 2)
        self.assertEqual(approval["current_step_name"], "导师初审")
        self.assertFalse(approval["can_review"])
        self.assertTrue(approval["is_requester"])
        self.assertFalse(created["data"][0]["can_review"])

        # 已有进行中的预约时，不允许删除流程或调整节点
        self.assertEqual(
            self.client.delete(f"/api/approval-flows/{flow_id}", headers=admin_headers).status_code,
            400,
        )
        changed_payload = {
            **flow_payload,
            "steps": [
                {"name": "导师初审", "approver_type": "role", "approver_role": "teacher"},
                {"name": "管理员终审", "approver_type": "role", "approver_role": "admin"},
            ],
        }
        self.assertEqual(
            self.client.put(f"/api/approval-flows/{lab_id}", headers=admin_headers, json=changed_payload).status_code,
            400,
        )

        # 非当前节点审批人（同为教师）不能处理
        teacher_b_list = self.client.get(
            "/api/reservations", headers=teacher_b_headers, params={"pending_for_me": "true"}
        ).json()["data"]
        self.assertFalse(any(item["id"] == reservation_id for item in teacher_b_list))
        self.assertEqual(
            self.client.put(
                f"/api/reservations/{reservation_id}/review",
                headers=teacher_b_headers,
                json={"status": "approved", "review_comment": "越权"},
            ).status_code,
            403,
        )

        # 第一节点通过后仍是待审核，第二节点由管理员处理
        teacher_a_all = self.client.get("/api/reservations", headers=teacher_a_headers).json()["data"]
        teacher_a_row = next(item for item in teacher_a_all if item["id"] == reservation_id)
        self.assertTrue(teacher_a_row["can_review"])
        first_approved = self.client.put(
            f"/api/reservations/{reservation_id}/review",
            headers=teacher_a_headers,
            json={"status": "approved", "review_comment": "同意使用"},
        )
        self.assertEqual(first_approved.status_code, 200)
        self.assertIn("管理员终审", first_approved.json()["message"])
        teacher_a_rows = self.client.get(
            "/api/reservations", headers=teacher_a_headers, params={"pending_for_me": "true"}
        ).json()["data"]
        self.assertFalse(any(item["id"] == reservation_id for item in teacher_a_rows))

        admin_rows = self.client.get(
            "/api/reservations", headers=admin_headers, params={"pending_for_me": "true"}
        ).json()["data"]
        admin_row = next(item for item in admin_rows if item["id"] == reservation_id)
        self.assertEqual(admin_row["status"], "pending")
        self.assertEqual(admin_row["approval"]["current_step"], 2)
        self.assertTrue(admin_row["approval"]["can_review"])

        with SessionLocal() as db:
            pending_notifications = (
                db.query(Notification)
                .filter(
                    Notification.type == "approval_pending",
                    Notification.related_id == reservation_id,
                    Notification.title == "待审批预约",
                )
                .all()
            )
            self.assertTrue(pending_notifications)

        final_approved = self.client.put(
            f"/api/reservations/{reservation_id}/review",
            headers=admin_headers,
            json={"status": "approved", "review_comment": "准予使用"},
        )
        self.assertEqual(final_approved.status_code, 200)
        self.assertIn("审核完成", final_approved.json()["message"])
        self.assertEqual(
            self.client.put(
                f"/api/reservations/{reservation_id}/review",
                headers=admin_headers,
                json={"status": "approved"},
            ).status_code,
            400,
        )

        # 审批进度详情：申请人可看节点与轨迹，其他学生不可见
        detail = self.client.get(f"/api/reservations/{reservation_id}/approval", headers=student_headers).json()
        self.assertEqual(detail["code"], 200)
        self.assertEqual([step["state"] for step in detail["data"]["steps"]], ["approved", "approved"])
        self.assertEqual(len(detail["data"]["decisions"]), 2)
        self.assertEqual(detail["data"]["decisions"][0]["approver_name"], "flow_teacher_a")
        self.assertFalse(detail["data"]["can_review"])

        self.client.post(
            "/api/auth/register",
            json={
                "username": "flow_other_student",
                "password": "FlowStu123456",
                "email": "flow_other_student@example.com",
            },
        )
        other_headers = self.login("flow_other_student", "FlowStu123456")
        self.assertEqual(
            self.client.get(f"/api/reservations/{reservation_id}/approval", headers=other_headers).status_code,
            403,
        )

        # 驳回路径：首节点拒绝即结束，未走到的第二节点标记为跳过
        second_date = today + timedelta(days=14)
        rejected = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": second_date.isoformat(),
                "start_time": "09:00:00",
                "end_time": "10:00:00",
                "purpose": "多级审批驳回单",
            },
        ).json()["data"][0]
        rejected_id = rejected["id"]
        rejected_result = self.client.put(
            f"/api/reservations/{rejected_id}/review",
            headers=teacher_a_headers,
            json={"status": "rejected", "review_comment": "用途不明确"},
        )
        self.assertEqual(rejected_result.status_code, 200)
        self.assertIn("流程结束", rejected_result.json()["message"])
        rejected_detail = self.client.get(f"/api/reservations/{rejected_id}/approval", headers=student_headers).json()
        self.assertEqual(rejected_detail["data"]["steps"][0]["state"], "rejected")
        self.assertEqual(rejected_detail["data"]["steps"][1]["state"], "skipped")

        # 首节点通过后学生改约：已产生的审批作废，从首节点重新开始
        third_date = today + timedelta(days=21)
        edited = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": third_date.isoformat(),
                "start_time": "09:00:00",
                "end_time": "10:00:00",
                "purpose": "可修改的预约",
            },
        ).json()["data"][0]
        edited_id = edited["id"]
        self.assertEqual(
            self.client.put(
                f"/api/reservations/{edited_id}/review",
                headers=teacher_a_headers,
                json={"status": "approved", "review_comment": "初审通过"},
            ).status_code,
            200,
        )
        updated = self.client.put(
            f"/api/reservations/{edited_id}",
            headers=student_headers,
            json={"purpose": "改约后的用途", "start_time": "10:00:00", "end_time": "11:00:00"},
        ).json()
        self.assertEqual(updated["code"], 200)
        self.assertIn("重新开始", updated["message"])
        self.assertEqual(updated["data"]["approval"]["current_step"], 1)
        reset_detail = self.client.get(f"/api/reservations/{edited_id}/approval", headers=student_headers).json()
        self.assertEqual(reset_detail["data"]["decisions"], [])
        self.assertEqual(reset_detail["data"]["steps"][0]["state"], "pending")

        # 停用后新预约不再进入流程；历史预约仍保留流程快照且不能删除流程
        disabled_payload = {**flow_payload, "enabled": False}
        disabled = self.client.put(f"/api/approval-flows/{lab_id}", headers=admin_headers, json=disabled_payload).json()
        self.assertEqual(disabled["code"], 200)
        self.assertFalse(disabled["data"]["enabled"])
        fourth_date = today + timedelta(days=28)
        plain = self.client.post(
            "/api/reservations",
            headers=student_headers,
            json={
                "lab_id": lab_id,
                "reservation_date": fourth_date.isoformat(),
                "start_time": "09:00:00",
                "end_time": "10:00:00",
                "purpose": "停用流程后的普通预约",
            },
        ).json()["data"][0]
        self.assertIsNone(plain["approval"])
        self.assertFalse(plain["can_review"])
        plain_row = next(
            item
            for item in self.client.get("/api/reservations", headers=admin_headers).json()["data"]
            if item["id"] == plain["id"]
        )
        self.assertTrue(plain_row["can_review"])
        self.assertEqual(
            self.client.delete(f"/api/approval-flows/{flow_id}", headers=admin_headers).status_code,
            400,
        )

        # 未被任何预约使用的新流程可以正常删除
        spare_lab_id = self.client.post(
            "/api/labs",
            headers=admin_headers,
            json={"name": "可删除流程实验室", "location": "实训楼 902", "capacity": 5},
        ).json()["data"]["id"]
        spare = self.client.put(
            f"/api/approval-flows/{spare_lab_id}",
            headers=admin_headers,
            json={
                "name": "备用流程",
                "steps": [{"name": "管理员审核", "approver_type": "role", "approver_role": "admin"}],
            },
        ).json()
        self.assertEqual(spare["code"], 200)
        self.assertEqual(
            self.client.delete(f"/api/approval-flows/{spare['data']['id']}", headers=admin_headers).status_code,
            200,
        )

        with SessionLocal() as db:
            stored = db.get(Reservation, reservation_id)
            self.assertEqual(stored.status, "approved")
            self.assertEqual(stored.approval_step_order, 2)

        # 本用例的预约会进入首页“实验室使用排行”Top8，清掉数据避免影响后续统计断言。
        from app.models.approval import ReservationApproval

        with SessionLocal() as db:
            reservation_ids = [reservation_id, rejected_id, edited_id, plain["id"]]
            db.query(ReservationApproval).filter(ReservationApproval.reservation_id.in_(reservation_ids)).delete(
                synchronize_session=False
            )
            db.query(Notification).filter(Notification.related_id.in_(reservation_ids)).delete(
                synchronize_session=False
            )
            db.query(Reservation).filter(Reservation.lab_id == lab_id).delete(synchronize_session=False)
            db.commit()


class _ScriptedChatClient:
    """按脚本吐流的假模型客户端：每次 create 调用依次消费一段预置分片。

    如果脚本被多调用一次，会直接抛断言错误——这正是回归“模型被重复调用”问题：
    修复前一次问答会同时命中图内 agent 和图外 run_model_tools 两条路径。
    """

    def __init__(self, rounds):
        self.rounds = list(rounds)
        self.create_calls = 0
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        self.create_calls += 1
        if not self.rounds:
            raise AssertionError("模型调用次数超出脚本预期：说明存在重复调用路径")
        return iter(self.rounds.pop(0))


def _content_round(text, chunk_size=4):
    """构造一轮纯文本流式返回。"""
    return [
        SimpleNamespace(
            choices=[SimpleNamespace(delta=SimpleNamespace(content=text[index : index + chunk_size], tool_calls=None))]
        )
        for index in range(0, len(text), chunk_size)
    ]


def _tool_call_round(name, arguments, call_id="call_1"):
    """构造一轮工具调用返回：首片带 id 和名称，后续片增量下发参数。"""
    return [
        SimpleNamespace(
            choices=[
                SimpleNamespace(
                    delta=SimpleNamespace(
                        content=None,
                        tool_calls=[
                            SimpleNamespace(
                                index=0,
                                id=call_id,
                                function=SimpleNamespace(name=name, arguments=""),
                            )
                        ],
                    )
                )
            ]
        ),
        SimpleNamespace(
            choices=[
                SimpleNamespace(
                    delta=SimpleNamespace(
                        content=None,
                        tool_calls=[
                            SimpleNamespace(
                                index=0,
                                id=None,
                                function=SimpleNamespace(name=None, arguments=arguments),
                            )
                        ],
                    )
                )
            ]
        ),
    ]


class AiStreamConvergenceTest(unittest.TestCase):
    """回归：工具事件与回答 token 必须只来自图内一条路径。

    背景：修复前模型工具调用在图内 agent 和图外 run_model_tools 各执行一遍，
    模型被调用三次、写工具可能被重复执行。现在统一为图内单一路径。
    本类字母序先于 ApiFlowTest 执行且共享同一临时数据库：收尾时只清理自己的
    Chroma 索引目录，共享数据库文件留给 ApiFlowTest 统一删除。
    """

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.chroma_directory = tempfile.TemporaryDirectory(prefix="lab_agent_chroma_sse_test_")
        cls.chroma_patch = patch("app.ai.knowledge.CHROMA_DIR", Path(cls.chroma_directory.name))
        cls.chroma_patch.start()
        registered = cls.client.post(
            "/api/auth/register",
            json={
                "username": "sse_user",
                "password": "Sse123456",
                "email": "sse-user@example.com",
                "full_name": "流式回归用户",
            },
        )
        if registered.json()["code"] != 200:
            raise RuntimeError(f"测试用户注册失败：{registered.text}")

    @classmethod
    def tearDownClass(cls):
        cls.chroma_patch.stop()
        SharedSystemClient.clear_system_cache()
        gc.collect()
        cls.chroma_directory.cleanup()

    def login_sse_user(self) -> dict[str, str]:
        response = self.client.post(
            "/api/auth/login",
            json={"username": "sse_user", "password": "Sse123456"},
        )
        self.assertEqual(response.json()["code"], 200)
        return {"Authorization": f"Bearer {response.json()['data']['token']}"}

    def test_keyed_graph_is_the_only_tool_path(self):
        """有 Key 且需要一次工具调用：模型最多调 2 次、工具只执行 1 次、事件完整。"""
        headers = self.login_sse_user()
        chat_client = _ScriptedChatClient(
            [
                _tool_call_round("get_today", "{}"),
                _content_round("今天是 2026 年 10 月 17 日。"),
            ]
        )
        with (
            patch("app.ai.agent.settings.OPENAI_API_KEY", "test-key"),
            patch("app.ai.agent.OpenAI", return_value=chat_client),
            patch("app.ai.agent.get_today", return_value="2026-10-17") as today_tool,
            patch("app.api.ai.stream_answer") as fallback_answer,
        ):
            response = self.client.post(
                "/api/ai/chat/stream",
                headers=headers,
                json={"message": "今天几号？", "history": []},
            )

        self.assertEqual(response.status_code, 200)
        text = response.text
        # 修复前这里会变成 3 次模型调用、工具执行 2 次。
        self.assertEqual(chat_client.create_calls, 2)
        today_tool.assert_called_once()
        # 工具事件来自图内透传：每个工具一对 start/end，标签可读。
        self.assertEqual(text.count("event: tool_start"), 1)
        self.assertEqual(text.count("event: tool_end"), 1)
        self.assertIn("确认日期", text)
        # 回答 token 由图内流式产出，不再走第二段生成路径。
        self.assertEqual(extract_streamed_tokens(text), "今天是 2026 年 10 月 17 日。")
        fallback_answer.assert_not_called()
        self.assertIn("event: done", text)
        # “已调用”汇总出现在回答 token 之前，保持旧版展示顺序。
        self.assertLess(text.index("已调用："), text.index("event: token"))

    def test_keyed_graph_answers_without_tools_in_one_call(self):
        """有 Key 但不需要工具：一次模型调用完成回答，不出现任何工具事件。"""
        headers = self.login_sse_user()
        chat_client = _ScriptedChatClient([_content_round("实验室周一至周五开放。")])
        with (
            patch("app.ai.agent.settings.OPENAI_API_KEY", "test-key"),
            patch("app.ai.agent.OpenAI", return_value=chat_client),
            patch("app.api.ai.stream_answer") as fallback_answer,
        ):
            response = self.client.post(
                "/api/ai/chat/stream",
                headers=headers,
                json={"message": "实验室什么时候开放？", "history": []},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(chat_client.create_calls, 1)
        self.assertEqual(extract_streamed_tokens(response.text), "实验室周一至周五开放。")
        fallback_answer.assert_not_called()
        self.assertNotIn("event: tool_start", response.text)

    def test_write_tool_runs_at_most_once_per_turn(self):
        """同一轮对话里创建预约最多真正执行一次，重复调用只回提示。"""
        events = []
        arguments = json.dumps(
            {
                "lab_id": 1,
                "reservation_date": "2099-01-01",
                "start_time": "09:00",
                "end_time": "10:00",
                "purpose": "设备调试",
            },
            ensure_ascii=False,
        )
        state = {
            "message": "确认预约",
            "user_id": 1,
            "db": None,
            "event_sink": events.append,
            "messages": [
                {
                    "role": "assistant",
                    "tool_calls": [
                        {
                            "id": "call_1",
                            "type": "function",
                            "function": {"name": "create_lab_reservation", "arguments": arguments},
                        }
                    ],
                }
            ],
        }
        with patch("app.ai.agent.create_lab_reservation", return_value="预约创建成功") as create_tool:
            first = execute_tools_node(state)
            self.assertEqual(create_tool.call_count, 1)
            self.assertTrue(first["write_tool_handled"])

            second_state = {
                **state,
                **first,
                "messages": [
                    {
                        "role": "assistant",
                        "tool_calls": [
                            {
                                "id": "call_2",
                                "type": "function",
                                "function": {"name": "create_lab_reservation", "arguments": arguments},
                            }
                        ],
                    }
                ],
            }
            second = execute_tools_node(second_state)
            self.assertEqual(create_tool.call_count, 1)
            self.assertIn("已经处理过一次预约请求", second["messages"][-1]["content"])

        self.assertEqual(
            [event["type"] for event in events],
            ["tool_start", "tool_end", "tool_start", "tool_end"],
        )

    def test_agent_write_tool_idempotency_and_audit(self):
        """写工具：同一用户+工具+参数的重复请求只执行一次，并写入审计记录。"""
        from app.models.agent_tool_log import AgentToolLog

        arguments = json.dumps(
            {
                "lab_id": 1,
                "reservation_date": "2099-01-02",
                "start_time": "10:00:00",
                "end_time": "11:00:00",
                "purpose": "幂等测试",
            },
            ensure_ascii=False,
        )
        tool_calls = [
            {
                "id": "call_idem",
                "type": "function",
                "function": {"name": "create_lab_reservation", "arguments": arguments},
            }
        ]

        with SessionLocal() as db:
            # 本测试类先于 ApiFlowTest 执行，不能依赖 admin 已存在，这里自建用户
            user = User(
                username="idem_user",
                password=hash_password("IdemUser123456"),
                email="idem_user@example.com",
                full_name="幂等测试用户",
                is_active=True,
                is_superuser=False,
            )
            db.add(user)
            db.commit()
            db.refresh(user)
            before = db.query(AgentToolLog).count()
            state = {
                "message": "确认预约",
                "user_id": user.id,
                "db": db,
                "event_sink": [].append,
                "messages": [{"role": "assistant", "tool_calls": tool_calls}],
            }
            with patch(
                "app.ai.agent.create_lab_reservation",
                return_value="✅ 预约成功！申请人：幂等测试用户，1条预约编号 1001，状态为待审核。",
            ) as create_tool:
                execute_tools_node(state)
                execute_tools_node(state)  # 相同请求重放，应命中幂等键
            self.assertEqual(create_tool.call_count, 1)

            logs = db.query(AgentToolLog).filter(AgentToolLog.tool_name == "create_lab_reservation").all()
            self.assertEqual(len(logs), 1)
            self.assertTrue(logs[0].success)
            self.assertIsNotNone(logs[0].idempotency_key)
            self.assertEqual(db.query(AgentToolLog).count(), before + 1)

    def test_read_tool_timeout_and_cancel_skip(self):
        """读工具硬超时；客户端断连后跳过尚未执行的工具。"""
        import threading
        import time
        from unittest.mock import patch

        from app.ai import agent as agent_module
        from app.config import settings

        def make_state(**extra):
            return {
                "message": "列出开放实验室",
                "user_id": 1,
                "db": None,
                "event_sink": lambda event: None,
                "messages": [
                    {
                        "role": "assistant",
                        "tool_calls": [
                            {
                                "id": "call_read",
                                "type": "function",
                                "function": {"name": "list_open_labs", "arguments": "{}"},
                            }
                        ],
                    }
                ],
                **extra,
            }

        with (
            patch.object(agent_module, "list_open_labs", side_effect=lambda db: time.sleep(0.5) or "实验室列表"),
            patch.object(settings, "TOOL_TIMEOUT_SECONDS", 0.05),
        ):
            timed_out = execute_tools_node(make_state())
        self.assertIn("超时", timed_out["messages"][-1]["content"])

        cancel_event = threading.Event()
        cancel_event.set()
        with patch.object(agent_module, "list_open_labs") as tool:
            cancelled = execute_tools_node(make_state(cancel_event=cancel_event))
        tool.assert_not_called()
        self.assertIn("取消", cancelled["messages"][-1]["content"])

    def test_model_tool_path_stays_single_source(self):
        """结构回归：重复的图外模型工具入口不得回流，端点也不再自建数据库会话。"""
        import inspect

        import app.ai.agent as agent_module
        import app.api.ai as ai_api_module

        self.assertFalse(hasattr(agent_module, "run_model_tools"))
        api_source = inspect.getsource(ai_api_module)
        self.assertNotIn("run_model_tools", api_source)
        self.assertNotIn("SessionLocal", api_source)


if __name__ == "__main__":
    unittest.main()
